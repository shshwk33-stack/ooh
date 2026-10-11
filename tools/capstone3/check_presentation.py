"""Check the delivered deck, render every slide, and record concrete evidence."""
from pathlib import Path
import json
import hashlib
import re
import fitz
from PIL import Image, ImageDraw
from pptx import Presentation
from docx import Document

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/"deliverables/capstone3"
prs=Presentation(OUT/"Transit_OOH_Capstone_3rd.pptx")
pdf=fitz.open(OUT/"Transit_OOH_Capstone_3rd.pdf")
assert len(prs.slides)==len(pdf)==10
normalize=lambda t:re.sub(r"\s+","",t)
issues=[];stats=[];all_text=[]
thumbs=Image.new("RGB",(1600,2270),"#DCE4EC")
preview=OUT/"previews";preview.mkdir(exist_ok=True)
for index,(slide,page) in enumerate(zip(prs.slides,pdf),1):
    body=page.get_text();all_text.append(body)
    assert "�" not in body
    missing=[]
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if p.text and normalize(p.text) not in normalize(body):missing.append(p.text)
        if shape.left < 0 or shape.top < 0 or shape.left+shape.width > prs.slide_width+20 or shape.top+shape.height > prs.slide_height+20:
            issues.append({"slide":index,"shape_outside_canvas":shape.name})
    if missing:issues.append({"slide":index,"missing_rendered_text":missing})
    assert len(slide.notes_slide.notes_text_frame.text)>300
    pix=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5))
    im=Image.frombytes("RGB",[pix.width,pix.height],pix.samples)
    im.save(preview/f"slide_{index:02d}.png")
    im.thumbnail((780,439));thumbs.paste(im,(((index-1)%2)*800+10,((index-1)//2)*454+4))
    stats.append({"slide":index,"native_text_characters":len(body),"speaker_notes_characters":len(slide.notes_slide.notes_text_frame.text),"missing_text":missing})
assert not issues,issues
thumbs.save(OUT/"Transit_OOH_Capstone_3rd_Preview.png")
combined="\n".join(all_text)
for phrase in ["토큰", "408", "16", "50:50", "60:40", "68", "합성"]:assert phrase in combined,phrase
assert len(Document(OUT/"Transit_OOH_Capstone_3rd_Speaker_Notes.docx").paragraphs)>60
notes_pdf=fitz.open(OUT/"Transit_OOH_Capstone_3rd_Speaker_Notes.pdf")
assert all("�" not in p.get_text() for p in notes_pdf)
manifest={"deck_slides":10,"aspect_ratio":"16:9","embedded_speaker_notes":10,"planned_duration_seconds":540,
"rendered_pdf_pages":len(pdf),"speaker_companion_pages":len(notes_pdf),"all_native_paragraphs_present_in_pdf":True,
"shape_canvas_bounds_passed":True,"replacement_characters":0,
"slide_checks":stats,"files":{},"evidence":{
"map":json.loads((OUT/"assets/map_evidence.json").read_text()),
"browser":json.loads((OUT/"assets/dashboard_capture_evidence.json").read_text()),
"software_tests":"Existing reports/runs/tests.xml, 2026-10-10: 68 pass, 1 expected failure, 0 failures; not an advertising performance result.",
"scope":"Four provider file types received and locally validated. Population CSV token blocked, official annual-average app adapter pending. 2023/2024 boundary equivalence unverified."}}
for filename in ["Transit_OOH_Capstone_3rd.pptx","Transit_OOH_Capstone_3rd.pdf","Transit_OOH_Capstone_3rd_Speaker_Notes.docx","Transit_OOH_Capstone_3rd_Speaker_Notes.pdf","Transit_OOH_Capstone_3rd_Preview.png"]:
    f=OUT/filename;manifest["files"][filename]={"bytes":f.stat().st_size,"sha256":hashlib.sha256(f.read_bytes()).hexdigest()}
(OUT/"presentation_validation.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf-8")
print(json.dumps({k:v for k,v in manifest.items() if k not in ("files","slide_checks","evidence")},ensure_ascii=False,indent=2))

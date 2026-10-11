"""Build a 10-slide, editable Korean progress deck and speaker companion."""
from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from docx import Document
from docx.shared import Inches as DInches, Pt as DPt, RGBColor as DColor
from docx.oxml import OxmlElement as DOxml
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "deliverables/capstone3"
ASSETS = OUT / "assets"
OUT.mkdir(parents=True, exist_ok=True)
FONT = "Noto Sans CJK KR"
NAVY="10283E"; INK="183249"; MUTED="62778A"; BLUE="2854D8"; TEAL="078893"
CYAN="48D8CA"; LIGHT="F3F6FA"; WHITE="FFFFFF"; LINE="D8E2EB"
PALE="E9F0FF"; MINT="E5F5F3"; AMBER="A86012"; CREAM="FFF3E3"; DARK2="19384F"
prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
prs.core_properties.title="Transit OOH | 캡스톤디자인2 3차 발표"
prs.core_properties.subject="데이터 확보와 분석 시스템 구현 현황"
prs.core_properties.author="Transit OOH 프로젝트팀"
prs.core_properties.keywords="정보통신공학, 캡스톤디자인2, 마포구, View-T, GPS, GIS"
slides=[]

def rgb(c):return RGBColor.from_string(c)
def box(s,x,y,w,h,fill=WHITE,stroke=None,radius=False,lw=.8):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
    if stroke:sh.line.color.rgb=rgb(stroke);sh.line.width=Pt(lw)
    else:sh.line.fill.background()
    sh._element.spPr.append(OxmlElement("a:effectLst"))
    if radius:
        try:sh.adjustments[0]=.12
        except Exception:pass
    return sh
def text(s,x,y,w,h,string,size=20,color=INK,bold=False,align=PP_ALIGN.LEFT,valign=MSO_ANCHOR.TOP):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=valign
    for i,line in enumerate(string.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align;p.space_before=Pt(0);p.space_after=Pt(4);p.line_spacing=1.12
        r=p.add_run();r.text=line;r.font.name=FONT;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=rgb(color)
        pr=r._r.get_or_add_rPr()
        ea=OxmlElement("a:ea");ea.set("typeface",FONT);pr.append(ea)
    return sh
def line(s,x1,y1,x2,y2,color=LINE,width=1.2,arrow=False,dash=False):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=rgb(color);sh.line.width=Pt(width)
    ln=sh._element.spPr.get_or_add_ln()
    if arrow:
        el=OxmlElement("a:tailEnd");el.set("type","triangle");el.set("w","sm");el.set("len","sm");ln.append(el)
    if dash:
        el=OxmlElement("a:prstDash");el.set("val","dash");ln.append(el)
    return sh
def circle(s,x,y,d,fill,stroke=None):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
    if stroke:sh.line.color.rgb=rgb(stroke)
    else:sh.line.fill.background()
    return sh
def pill(s,x,y,w,label,fill=MINT,color=TEAL,size=13):
    box(s,x,y,w,.37,fill,radius=True)
    text(s,x+.06,y+.055,w-.12,.27,label,size,color,True,PP_ALIGN.CENTER)
def picture(s,path,x,y,w=None,h=None):
    return s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w) if w else None,height=Inches(h) if h else None)
def picture_contain(s,path,x,y,w,h):
    with Image.open(path) as im: ratio=im.width/im.height
    actual_w=min(w,h*ratio);actual_h=actual_w/ratio
    return picture(s,path,x+(w-actual_w)/2,y+(h-actual_h)/2,w=actual_w)
def page(n,category,title,subtitle="",dark=False):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(NAVY if dark else LIGHT)
    color=WHITE if dark else INK
    text(s,.62,.36,10,.28,f"TRANSIT OOH  /  CAPSTONE DESIGN II  /  {category}",10.5,CYAN if dark else TEAL,True)
    text(s,.62,.93,14.7,.65,title,31,color,True)
    if subtitle:text(s,.65,1.69,14.6,.43,subtitle,16.5,"B9CDD9" if dark else MUTED)
    line(s,.64,8.48,15.34,8.48,"335169" if dark else LINE,.7)
    text(s,.65,8.60,13.8,.22,"정보통신공학과 · 캡스톤디자인2 · 3차 진행 보고",9,"A4BDC9" if dark else MUTED)
    text(s,14.35,8.54,1,.29,f"{n:02d} / 10",11,"A4BDC9" if dark else MUTED,False,PP_ALIGN.RIGHT)
    return s
def source(s,string,dark=False):
    text(s,.66,8.12,14.65,.26,string,9.1,"9EB8C7" if dark else MUTED)
def notes(s,title,body,evidence,seconds):
    content=f"{title}\n권장 발표시간: 약 {seconds}초\n\n{body}\n\n근거 / 설명 범위\n{evidence}"
    s.notes_slide.notes_text_frame.text=content
    slides.append({"title":title,"seconds":seconds,"script":body,"evidence":evidence})
def small_icon(s,kind,x,y,color=TEAL):
    if kind=="file":
        box(s,x,y,.43,.53,WHITE,color)
        for dy in [.17,.28,.39]:line(s,x+.08,y+dy,x+.34,y+dy,color,1)
    elif kind=="map":
        for xx,yy in [(.03,.29),(.21,.06),(.45,.37)]:circle(s,x+xx,y+yy,.11,color)
        line(s,x+.08,y+.34,x+.26,y+.11,color,2);line(s,x+.26,y+.11,x+.50,y+.42,color,2)
    elif kind=="check":
        line(s,x+.04,y+.29,x+.19,y+.45,color,2.5);line(s,x+.19,y+.45,x+.52,y+.08,color,2.5)
    elif kind=="chart":
        for xx,hh in [(.02,.22),(.20,.42),(.38,.32)]:box(s,x+xx,y+.52-hh,.11,hh,color)

# 01 — editorial cover, using the real provider network as a geographic motif.
s=page(1,"PROGRESS REPORT 03","",dark=True)
box(s,8.57,.82,6.77,6.76,DARK2,radius=True)
picture(s,ASSETS/"mapo_cover_network.png",8.86,1.30,w=6.18)
text(s,9.03,1.08,5.5,.36,"MAPO-GU  /  ROAD NETWORK",13,CYAN,True)
text(s,9.03,6.71,5.8,.38,"2024년 주요도로망 · 마포구 408개 구간",15,WHITE,True)
text(s,9.03,7.14,5.7,.28,"확보한 도로 형상 표시 · 배송차량 운행궤적 아님",11,"B9CDD9")
text(s,.67,1.46,7.8,.82,"Transit OOH",46,WHITE,True)
text(s,.7,2.65,7.7,1.50,"데이터를 모으고,\n분석 시스템으로 연결하다",31,WHITE,True)
text(s,.73,4.54,7.25,.85,"배송차량 이동정보를 활용한\n상대적 위치가치 분석 시스템",21,"C5D8E5")
pill(s,.75,5.88,2.47,"3차 발표 · 진행 보고",fill="28485F",color=CYAN)
text(s,.77,6.47,7.20,.70,"데이터 확보 방법  /  확보 범위  /  구현 현황",17,WHITE)
text(s,.76,7.51,7.0,.30,"정보통신공학과  ·  캡스톤디자인2",15,"B9CDD9",True)
source(s,"자료 확인 기준 2026.10.11 | 지도 출처: 한국교통연구원 KTDB View-T, 2024년 서울 주요도로망",True)
notes(s,"01. 이번 발표의 초점", "안녕하십니까. Transit OOH 3차 진행 상황을 발표하겠습니다. 저희는 배송차량의 이동정보와 지역 특성을 함께 사용해, 어떤 장소를 우선 검토할지 돕는 분석 시스템을 만들고 있습니다. 이번 발표에서는 실제 데이터를 어디에서 구했고, 현재 어떤 파일을 확보했으며, 이를 처리할 프로그램을 어디까지 구현했는지 설명드리겠습니다. 오른쪽은 이번에 받은 마포구 도로망입니다. 차량이 실제 주행한 GPS 경로를 그린 것은 아닙니다.", "공식 도로망 파일의 공급자 마포구 코드로 408개 링크를 추출. 자료의 정의와 제약은 VIEWT_COLLECTION_STATUS.md 및 network_validation_summary.json 참조.",35)

# 02 — progress, not an inflated final-result claim.
s=page(2,"01  연구의 흐름","아이디어 검증에서, 실제 자료 연결 준비로","기존 연구 방향은 유지하고, 이번에는 데이터와 시스템의 연결 기반을 구체화했다.")
cards=[(.66,"1차 발표","왜 필요한가","배송 이동정보의 추가 활용\n연구 질문과 시스템 방향 제시",WHITE,TEAL),(5.61,"2차 발표","구현 가능한가","합성 GPS로 초기 분석 실험\n정차·점수 계산·지도 표현 확인",WHITE,BLUE),(10.56,"3차 발표","실제 자료를 어떻게 붙일까","공식 파일 확보와 구조 확인\n분석 모듈·웹 화면 구현 현황",PALE,BLUE)]
for x,k,t,b,fill,c in cards:
    box(s,x,2.40,4.77,2.53,fill,radius=True)
    pill(s,x+.24,2.65,1.28,k,fill=MINT if k.startswith("1") else PALE,color=c)
    text(s,x+.25,3.20,4.25,.48,t,23,c,True)
    text(s,x+.25,3.91,4.28,.76,b,17,INK)
line(s,5.02,3.63,5.48,3.63,BLUE,2,True);line(s,9.97,3.63,10.43,3.63,BLUE,2,True)
box(s,.66,5.32,14.67,1.52,NAVY,radius=True)
text(s,.95,5.57,2.10,.30,"유지하는 연구 질문",14,CYAN,True)
text(s,3.24,5.58,11.55,.85,"차량의 이동과 공간의 특성을 함께 보면,\n어떤 장소를 우선 검토할 수 있을까?",23,WHITE,True)
text(s,.84,7.19,14.25,.54,"오늘의 범위   자료 확보 → 입력 검증 → 시스템 구현     |     후속 발표   실제 자료 적용 → 비교·검증",16,MUTED)
source(s,"근거: 1차 발표 PDF · 2차 발표 PPTX/대본 · 현재 소스 및 진행 기록 | 순위 변화만으로 광고효과 향상을 주장하지 않음")
notes(s,"02. 1·2차 발표와 이어지는 지점", "1차 발표에서는 배송 과정에서 쌓이는 데이터를 새로운 의사결정에 활용할 수 있는지 질문했습니다. 2차 발표에서는 합성 GPS로 정차를 찾고 점수를 계산한 뒤 지도에 표현하는 초기 실험을 했습니다. 이번 3차는 그 구조에 실제 자료를 연결하기 위한 준비 단계입니다. 그래서 새로운 최종 성능을 제시하기보다, 공식 파일을 확보한 과정과 구현 현황을 보여드리겠습니다. 연구 질문은 그대로입니다. 이동정보를 공간정보에 더했을 때 후보지를 어떻게 다르게 검토할 수 있는지를 확인하는 것입니다.", "기존 발표자료를 SOURCE_AUDIT.md와 대조. 현재 성과를 원본 2차 PoC에 소급하여 기재하지 않음.",45)

# 03 — official acquisition workflow.
s=page(3,"02  데이터 확보 방법","공식 사이트에서 지도와 지표를 나누어 확보","한국교통연구원 View-T에서 자료 종류별 경로를 확인하고, 내려받은 원본 파일을 보존했다.")
stages=[("01","공식 경로 확인","View-T 메뉴·자료 설명"),("02","조건 선택","연도·지역·요일·도로"),("03","파일 다운로드","지도 ZIP · 지표 Excel"),("04","내용 검증","개수·번호·단위·누락")]
for i,(num,title,desc) in enumerate(stages):
    x=.67+i*3.73
    circle(s,x,2.49,.52,BLUE);text(s,x,2.57,.52,.29,num,14,WHITE,True,PP_ALIGN.CENTER)
    text(s,x+.70,2.44,2.75,.43,title,21,INK,True);text(s,x+.70,3.02,2.78,.42,desc,14.5,MUTED)
    if i<3:line(s,x+3.35,2.76,x+3.61,2.76,BLUE,1.6,True)
box(s,.67,3.81,14.67,1.67,WHITE,radius=True)
small_icon(s,"file",.99,4.21)
text(s,1.72,4.06,4.45,.40,"네트워크 · 차량통행지표",22,TEAL,True)
text(s,1.74,4.67,4.40,.44,"도로망·행정경계·교통량·속도",16,MUTED)
line(s,6.58,4.66,7.13,4.66,TEAL,2,True)
text(s,7.46,4.07,5.55,.43,"서울 → 마포구 / 2024년 평일",21,INK,True)
text(s,7.48,4.68,5.65,.43,"도로 지표는 연평균 일 자료로 확보",16,MUTED)
pill(s,13.33,4.42,1.59,"4종 확보")
box(s,.67,5.73,14.67,1.67,CREAM,radius=True)
small_icon(s,"chart",.99,6.13,AMBER)
text(s,1.72,5.98,4.76,.40,"Expert · 인구 이동 분석",22,AMBER,True)
text(s,1.74,6.59,4.64,.42,"전국 출발 → 마포구 유입 조건 설정",16,MUTED)
line(s,6.58,6.59,7.13,6.59,AMBER,2,True)
text(s,7.46,5.99,5.65,.43,"CSV 다운로드 단계에서 중단",21,INK,True)
text(s,7.48,6.60,5.7,.43,"토큰키 필요 · 인구 파일은 아직 없음",16,MUTED)
pill(s,13.33,6.34,1.59,"인증 대기",CREAM,AMBER)
source(s,"출처: viewt.ktdb.go.kr · 공식 네트워크/차량통행지표/Expert 화면 | 행정경계는 다운로드 가능한 2023년 자료 사용")
notes(s,"03. 데이터를 확보한 방법", "데이터는 한국교통연구원의 View-T 공식 사이트에서 확보했습니다. 먼저 자료 설명과 메뉴를 확인한 뒤 연도, 지역, 요일을 맞춰 파일을 받았습니다. 도로망과 행정경계는 지도 ZIP 파일로, 교통량과 평균속도는 Excel로 받았습니다. 도로 지표의 조건은 2024년 마포구 평일이고, 실제 파일에는 연평균 하루 값이 들어 있습니다. 인구 이동 분석은 별도 Expert 메뉴에서 마포구로 들어오는 조건을 설정했습니다. 다만 CSV를 받으려면 토큰키가 필요해서 현재는 다운로드 전 단계에서 멈춰 있습니다. API 자동 수집을 완료한 상태는 아닙니다.", "공식 페이지 및 사용자 다운로드 원본의 조건 행 확인. 연평균 일 자료는 날짜별 원시 관측값이 아님. 인구 분석의 설정값은 데이터 확보를 뜻하지 않음.",55)

# 04 — actual spatial scope.
s=page(4,"03  연구 대상 지역","연구 범위는 서울 마포구 한 곳으로 좁혔다","한 자치구 안에서 자료의 범위와 연결 관계를 먼저 확인하는 사례 연구로 시작한다.")
box(s,.67,2.26,9.39,5.18,"F0F5F7",radius=True)
picture_contain(s,ASSETS/"mapo_reference_map.png",.78,2.59,9.16,4.74)
pill(s,.92,2.48,2.72,"확보 파일로 그린 참고 지도",WHITE,TEAL,12)
box(s,10.35,2.27,4.98,1.23,WHITE,radius=True)
text(s,10.60,2.45,1.6,.66,"16",39,TEAL,True)
text(s,12.0,2.50,2.98,.57,"개 동 경계\n2023년 공간자료",17,INK)
box(s,10.35,3.67,4.98,1.23,WHITE,radius=True)
text(s,10.60,3.85,1.72,.66,"408",39,BLUE,True)
text(s,12.24,3.90,2.82,.62,"개 도로 구간\n2024년 주요도로망",16.5,INK)
text(s,10.46,5.22,4.69,.38,"분석할 운행 형태: 밀크런",20,INK,True)
text(s,10.48,5.83,4.61,.83,"여러 거점을 정해진 순서로\n반복 방문하는 배송 방식",18,MUTED)
for i in range(3):
    circle(s,10.66+i*1.69,6.87,.31,TEAL)
    if i<2:line(s,11.01+i*1.69,7.025,12.22+i*1.69,7.025,TEAL,1.6,True)
text(s,10.48,7.38,4.72,.27,"운행 형태 설명용 · 실제 배송 GPS 미확보",11,AMBER)
source(s,"지도 출처: View-T 2023년 마포구 경계 + 2024년 주요도로망 | 연도 차이가 있는 참고 중첩이며 경계 동일성은 추가 확인 필요")
notes(s,"04. 마포구를 기준으로 범위 설정", "전국을 한꺼번에 분석하기보다 마포구 한 곳을 예시 지역으로 정했습니다. 지도는 실제 확보한 파일로 그렸고, 16개 동의 경계와 공급자가 마포구로 분류한 408개 도로 구간을 보여줍니다. 연구 대상은 여러 거점을 반복적으로 도는 밀크런형 배송차량입니다. 아래 연결 그림은 이 운행 방식을 설명하는 그림이며, 관측한 차량 경로는 아닙니다. 경계는 2023년, 도로는 2024년 자료라는 차이도 있습니다. 따라서 지금 지도는 자료 범위를 확인하는 용도로 사용하고, 동별 점수를 계산하기 전에 경계 버전이 일치하는지 확인하겠습니다.", "마포구 도로 수는 지리적 클리핑 결과가 아니라 공급자 sigungu_id=11140 기준. 경계를 통과하는 모든 도로 수와 다를 수 있음. View-T 주요도로망은 모든 골목·배송 접근도로를 포함하지 않음.",55)

# 05 — source inventory, compact but report-ready.
s=page(5,"04  확보 현황","확보한 네 종류의 자료와, 아직 필요한 자료","파일을 받은 상태와 프로그램에 연결한 상태를 구분하여 관리하고 있다.")
cols=[.68,3.91,6.83,11.33];widths=[3.23,2.92,4.50,4.0]
headers=["자료","기준 / 단위","확인한 범위","현재 상태"]
for x,w,h in zip(cols,widths,headers):
    box(s,x,2.28,w,.49,NAVY);text(s,x+.16,2.39,w-.30,.28,h,15,WHITE,True)
rows=[
    ("행정경계","2023년 / 동","마포구 16개 동","확보 · 형상 확인",True),
    ("도로 네트워크","2024년 / 도로 구간","서울 8,582개 중 마포구 408개","확보 · 번호 확인",True),
    ("추정교통량","2024년 평일 / 대/일","408개 구간 · 전체/승용차/버스/트럭","확보 · 연결 확인",True),
    ("평균속도","2024년 평일 / km/h","408개 구간 · 전체 차량 평균","확보 · 연결 확인",True),
    ("인구 기종점 통행량","2024년 조건 설정","마포구 유입 분석까지 진행","미다운로드 · 토큰 필요",False),
    ("차량 대표 경로","동일 연도 자료 필요","경로 분석 도구·공식 문서 확인","미수집",False),
    ("배송 GPS · 동기간 시설정보","실제 운행 / 후보 구역","이동·정차·방문 및 주변 시설 검증용","추가 확보 필요",False),
]
for i,row in enumerate(rows):
    yy=2.77+i*.58;fill=WHITE if i%2==0 else "ECF1F6"
    for x,w,v in zip(cols,widths,row[:4]):
        box(s,x,yy,w,.58,fill)
        text(s,x+.16,yy+.145,w-.27,.35,v,14.5,TEAL if x==cols[-1] and row[4] else (AMBER if x==cols[-1] else INK),x in [cols[0],cols[-1]])
box(s,.68,7.09,14.65,.61,CREAM,radius=True)
text(s,.91,7.24,14.2,.35,"주의  교통량은 추정값, 속도는 집계값이다. 연평균 값을 실제 날짜별 관측이나 배송차량 GPS로 취급하지 않는다.",15,AMBER,True)
source(s,"근거: 업로드한 지도 ZIP 2종·교통량/속도 Excel 2종의 파일 검사 | 확보 자료의 공식 형식 변환·대시보드 연결은 후속 작업")
notes(s,"05. 어디까지 확보했는가", "현재 확보한 자료는 행정경계, 도로 네트워크, 추정교통량, 평균속도의 네 종류입니다. 교통량과 속도는 각각 마포구 408개 도로 구간에 대응합니다. 교통량에는 전체 차량과 승용차, 버스, 트럭이 나누어져 있습니다. 다만 이 값은 연평균 하루의 추정 통행량이고, 개별 차량을 추적한 기록은 아닙니다. 속도도 전체 차량의 평균이므로 배송차량의 정차시간으로 해석할 수 없습니다. 인구 CSV, 차량 대표 경로, 실제 배송 GPS와 분석 기간에 맞는 시설정보는 추가 확보가 필요합니다. 다른 공개 자료로 수행한 별도 분석을 마포구 자료로 섞어 넣지는 않았습니다.", "교통량 대/일과 평균속도 km/h는 서로 다른 단위. POI는 시설·장소 정보이며 기존 서울 일부 OSM 자료와 동기간 마포구용 완전한 자료는 구분.",60)

# 06 — data alignment evidence.
s=page(6,"05  입력 자료 검증","도로 번호를 기준으로, 세 파일이 연결되는지 확인했다","서로 같은 위치를 뜻하는지 먼저 확인해야 교통량과 속도를 지도에 붙일 수 있다.")
for i,(label,sub,col) in enumerate([("도로망","지도상의 위치와 도로 번호",TEAL),("교통량","같은 도로의 연평균 통행량",BLUE),("평균속도","같은 도로의 집계 속도",TEAL)]):
    x=.70+i*5.02
    box(s,x,2.30,4.63,1.48,WHITE,radius=True)
    text(s,x+.23,2.50,2.8,.4,label,22,col,True)
    text(s,x+2.89,2.48,1.45,.60,"408",32,col,True,PP_ALIGN.RIGHT)
    text(s,x+.23,3.14,4.19,.34,sub,15,MUTED)
    if i<2:line(s,x+4.69,3.01,x+4.97,3.01,BLUE,1.6,True)
box(s,.70,4.02,8.25,1.57,NAVY,radius=True)
text(s,1.01,4.24,4.15,.69,"408 / 408",37,WHITE,True)
text(s,5.13,4.34,3.40,.75,"도로 구간 번호 일치\n세 파일의 동일 대상 확인",18,CYAN,True)
box(s,9.21,4.02,6.10,1.57,MINT,radius=True)
text(s,9.54,4.25,5.40,.40,"중복 번호 0 · 누락 셀 0",23,TEAL,True)
text(s,9.55,4.89,5.35,.40,"교통량·속도 Excel의 검사 결과",16,MUTED)
checks=[("위치 기준","구역 번호와 좌표 기준 확인"),("시간 기준","2024년 · 평일 · 연평균 일 확인"),("의미 기준","대/일과 km/h를 따로 유지")]
for i,(title,desc) in enumerate(checks):
    x=.86+i*4.94;small_icon(s,"check",x,6.00,TEAL)
    text(s,x+.72,5.95,3.89,.42,title,21,INK,True)
    text(s,x+.73,6.61,3.96,.55,desc,16,MUTED)
text(s,.87,7.47,14.34,.34,"완료: 로컬 파일 대조     다음: 공식 연평균 자료에 맞는 변환 기능과 대시보드 연결",16,AMBER,True)
source(s,"근거: traffic_volume_validation.json · average_speed_validation.json | 408개는 확보한 주요도로 구간 수이며 차량 대수·후보지 수가 아님")
notes(s,"06. 파일을 검증한 방법", "파일을 받았다고 바로 합치는 것은 아닙니다. 먼저 도로망, 교통량, 속도가 같은 도로를 가리키는지 도로 구간 번호로 대조했습니다. 세 자료에서 마포구 408개 번호가 모두 일치했고, 교통량과 속도 표에는 중복 번호나 누락된 셀이 없었습니다. 이는 파일을 서로 연결할 수 있는 기초 조건을 확인한 결과입니다. 추가로 위치를 표현하는 좌표 기준, 자료의 연도와 평일 조건, 교통량과 속도의 단위를 확인했습니다. 다만 현재 앱은 날짜별 표준 입력을 처리하므로, 이번 연평균 파일을 안전하게 연결하려면 집계 방식에 맞는 별도 변환 기능이 필요합니다.", "정합성 검사는 로컬 업로드 파일에 한정. 2023년 경계와 2024년 도로 버전 동일성 검증 완료를 뜻하지 않음. XLSX A1:A1 잘못된 dimension은 reset_dimensions 후 전체 행을 확인. 소수점 형태 숫자 도로 ID는 정수성 확인 후 문자열로 처리.",60)

# 07 — implemented software architecture; two branches stay distinct.
s=page(7,"06  시스템 구현","입력 → 분석 → 지도, 기능을 나누어 구현했다","Python 분석 엔진과 Streamlit 웹 화면을 연결하고, 자료의 종류에 따라 처리 경로를 구분했다.",True)
labels=[("입력",.89),("검사",4.48),("정보 추출",8.06),("결과",11.85)]
for t,x in labels:text(s,x,2.31,3.0,.36,t,14,CYAN,True)
top=[(.72,"GPS 좌표·시간","위도·경도·관측 시각"),(4.31,"오류점·시간간격 검사","이상 이동·긴 공백 확인"),(7.89,"이동·정차·방문 집계","실제 경과시간으로 계산")]
for x,t,d in top:
    box(s,x,2.93,3.25,1.30,DARK2,radius=True)
    text(s,x+.22,3.12,2.83,.44,t,19,WHITE,True)
    text(s,x+.22,3.72,2.87,.31,d,14,"B9CDD9")
for x in [3.99,7.57]:line(s,x,3.57,x+.26,3.57,CYAN,1.8,True)
bot=[(.72,"교통·인구·시설 표","CSV·Excel 표준 형식"),(4.31,"지역·연도·단위 검사","조건이 다른 입력을 확인"),(7.89,"공간 번호로 자료 연결","같은 후보 구역으로 정리")]
for x,t,d in bot:
    box(s,x,4.84,3.25,1.30,DARK2,radius=True)
    text(s,x+.22,5.03,2.86,.44,t,18.5,WHITE,True)
    text(s,x+.22,5.64,2.87,.31,d,14,"B9CDD9")
for x in [3.99,7.57]:line(s,x,5.49,x+.26,5.49,CYAN,1.8,True)
box(s,11.68,2.93,3.60,3.21,"224D66",radius=True)
text(s,12.0,3.21,2.97,.85,"점수·순위 계산\nGIS 대시보드",24,WHITE,True)
line(s,12.0,4.53,14.94,4.53,"48778A",1)
text(s,12.0,4.83,3.0,.88,"가중치 조절 · 모델 비교\n후보 지도 · CSV 저장",16,CYAN)
line(s,11.16,3.57,11.61,3.57,CYAN,1.8,True);line(s,11.16,5.49,11.61,5.49,CYAN,1.8,True)
text(s,.87,4.37,10.25,.29,"GPS 처리 기능: 합성 시나리오로 구현·시험 · 실제 배송 GPS는 미확보",13,"B9CDD9")
text(s,.87,6.34,10.58,.52,"표준 형식 입력은 구현·시험 완료. 받은 공식 연평균 파일은 추가 변환이 필요하다.",14,"F3C78D")
line(s,.88,7.09,15.12,7.09,"335169",1)
text(s,.91,7.34,14.05,.36,"Python·Pandas  데이터 처리     |     GeoPandas·Shapely  공간 처리     |     Streamlit  웹 화면",16,"B9CDD9")
source(s,"구현 근거: src/transit_ooh/ · app/streamlit_app.py · app/viewt_page.py | View-T 집계자료만으로 GPS 정차·반복 방문을 측정하지 않음",True)
notes(s,"07. 정보통신공학 관점의 구현", "시스템은 크게 입력, 검사, 정보 추출, 결과 화면의 네 부분으로 나누었습니다. GPS가 들어오는 경로에서는 잘못된 위치와 긴 시간 공백을 확인하고, 이동시간과 정차시간, 방문 횟수를 계산합니다. 정차시간은 단순히 GPS 점 개수를 세는 방식 대신 관측 시각의 실제 차이를 사용합니다. 표 형태의 교통·인구 자료는 별도 경로에서 지역, 연도, 단위와 공간 번호를 확인합니다. 두 경로는 자료의 의미를 유지한 상태에서 점수 계산과 화면 출력으로 이어집니다. Python으로 계산하고 공간 처리 도구를 이용해 위치를 연결하며, Streamlit으로 사용자 화면을 구현했습니다. 현재 GPS 기능과 표준 파일 흐름은 합성 입력으로 시험했습니다.", "분석 모듈: cleaning, stops, features, scoring, experiments, viewt. GPS 점수와 집계 통행량 점수는 별도 구현. 집계자료의 공식 연평균 어댑터는 미완료. 정차 목적이 배송인지 신호대기인지는 GPS만으로 확정 불가.",60)

# 08 — easily understandable model and editable weight diagram.
s=page(8,"07  위치가치 계산 방식","이동 정보와 공간 정보의 비중을 사용자가 조절한다","같은 후보들을 같은 기준으로 점수화한 뒤, 사용 목적에 따라 두 정보의 비중을 바꾼다.")
box(s,.70,2.28,7.12,1.61,PALE,radius=True)
text(s,.96,2.47,6.54,.42,"M  이동 점수 · Mobility",24,BLUE,True)
text(s,.98,3.10,6.56,.65,"GPS: 이동시간·정차시간·방문 횟수\n집계자료 모드: 차량 통행량을 별도 지표로 사용",16,INK)
box(s,8.05,2.28,7.25,1.61,MINT,radius=True)
text(s,8.31,2.47,6.62,.42,"S  공간 점수 · Spatial",24,TEAL,True)
text(s,8.33,3.10,6.54,.65,"주변 인구 통행 규모와 시설 수를 반영\n인구 통행량을 실제 광고 노출 인원으로 보지 않음",16,INK)
text(s,.90,4.19,14.2,.49,"위치가치 = 이동 점수 × 비중 + 공간 점수 × 나머지 비중",25,INK,True,PP_ALIGN.CENTER)
box(s,1.04,5.02,6.96,.58,BLUE,radius=False);box(s,8.0,5.02,6.96,.58,TEAL)
text(s,1.21,5.14,6.53,.30,"이동 50%",18,WHITE,True,PP_ALIGN.CENTER)
text(s,8.18,5.14,6.51,.30,"공간 50%",18,WHITE,True,PP_ALIGN.CENTER)
circle(s,7.86,4.91,.29,WHITE,BLUE)
text(s,1.02,5.77,13.95,.36,"기본 50:50   ·   기존 60:40과 비교   ·   이동 비중 0~100% 조절",17,MUTED,False,PP_ALIGN.CENTER)
for i,(letter,title,c,fill) in enumerate([("A","공간만 평가",TEAL,MINT),("B","이동만 평가",BLUE,PALE),("C","두 정보를 결합",BLUE,PALE)]):
    x=.84+i*4.96;box(s,x,6.47,4.51,.71,WHITE,radius=True)
    circle(s,x+.19,6.64,.36,c);text(s,x+.19,6.69,.36,.23,letter,13,WHITE,True,PP_ALIGN.CENTER)
    text(s,x+.78,6.64,3.44,.37,title,20,INK,True)
text(s,.98,7.48,14.0,.33,"50:50과 60:40 모두 비교를 위한 설정값이다. 실제 성능으로 검증된 최적 비율은 아직 아니다.",16,AMBER,True,PP_ALIGN.CENTER)
source(s,"구현 근거: scoring.py · viewt.py | 각 입력은 0~1로 환산; GPS 이동 점수는 3개 지표 평균, 공간 점수의 70:30도 초기 가정")
notes(s,"08. 점수와 가중치의 의미", "최종 점수는 이동 점수와 공간 점수를 섞어서 계산합니다. 이동 점수는 GPS가 있을 때 이동시간, 정차시간, 방문 횟수를 사용합니다. 교통 집계자료를 쓰는 모드에서는 차량 통행량을 별도의 지표로 사용하며, GPS 지표와 같은 것으로 보지 않습니다. 공간 점수는 인구 통행 규모와 주변 시설정보를 반영합니다. 크기와 단위가 다른 값은 먼저 0에서 1 사이로 맞춥니다. 기본 비율은 이동과 공간을 반반 반영하는 50대50입니다. 사용자가 비중을 바꾸면 지도와 순위도 다시 계산됩니다. 공간만 보는 A, 이동만 보는 B, 함께 보는 C를 비교할 수 있습니다. 기존 60대40도 비교 대상인 초기 가정이며 최적값은 아닙니다.", "scoring.py: M=mean(z(moving_seconds),z(stop_seconds),z(visit_frequency)); S=.7z(population)+.3z(poi_count). viewt.py 집계 모드: M=z(vehicle_flow). I=alpha*M+(1-alpha)*S. 후보별 정규화 조건과 결측 검사는 구현 계약에 따름. 실제 마포구 점수 산출 완료를 의미하지 않음.",60)

# 09 — real app screenshot, explicit synthetic label, tests are software evidence.
s=page(9,"08  웹 대시보드 구현","사용자 조작과 결과 조회까지 웹 화면으로 연결했다","실행한 화면에서 자치구·시간대·가중치를 바꾸고, 지도와 후보 순위를 확인할 수 있다.")
box(s,.70,2.28,9.55,5.68,WHITE,LINE,radius=True)
pill(s,.91,2.42,3.65,"실제 실행 화면 · 입력은 합성 데이터",CREAM,AMBER,12)
picture_contain(s,ASSETS/"dashboard_viewport.png",.86,2.84,9.22,4.94)
features=[("01","조건 설정","자치구 · 파일 · 날짜/시간대\n가중치 0~100% 조절"),("02","결과 확인","지도 · 후보 순위 · 상위 5/10개\n가상 순회·모델별 비교"),("03","검토와 저장","가중치 변화 · 선정 안정성\n결과표 CSV 다운로드")]
for i,(n,t,d) in enumerate(features):
    yy=2.38+i*1.33;circle(s,10.66,yy,.40,BLUE);text(s,10.66,yy+.075,.4,.24,n,12,WHITE,True,PP_ALIGN.CENTER)
    text(s,11.24,yy-.02,3.92,.38,t,21,INK,True)
    text(s,11.26,yy+.49,3.9,.72,d,16,MUTED)
box(s,10.63,6.49,4.63,1.13,MINT,radius=True)
text(s,10.86,6.65,1.18,.62,"68",35,TEAL,True)
text(s,12.20,6.67,2.84,.70,"자동검사 통과\n실패 0 · 예상 실패 1",16,TEAL,True)
source(s,"화면: 현재 Streamlit 앱 직접 실행·촬영(합성 입력, 배경지도 OFF) | 검사: 2026.10.10 pytest, 기존 PoC 오류 1건은 예상 실패로 보존")
notes(s,"09. 대시보드에서 가능한 일", "이 화면은 실제 실행한 웹 프로그램을 촬영한 것입니다. 화면에 보이는 후보 16개와 지도상의 순회선은 기능 확인을 위한 합성 데이터입니다. 앞에서 보여드린 마포구의 실제 16개 동 경계와는 구분해야 합니다. 사용자는 자치구와 입력 파일, 날짜와 시간대, 이동정보 비중을 선택할 수 있습니다. 결과 화면에서는 지도와 후보 순위, 상위 후보, 모델 비교와 가중치 변화에 따른 안정성을 확인하고 표를 저장할 수 있습니다. 기존 자동검사는 68개가 통과했고 실패는 없었습니다. 과거 코드의 알려진 오류 한 건은 예상 실패로 남겨 재현합니다. 이 검사는 프로그램의 처리 기능을 확인한 것이며 실제 광고효과를 검증한 수치는 아닙니다.", "스크린샷: 실제 Chromium의 실행 화면, 합성 경고 유지, 외부 배경지도 미사용. reports/runs/tests.xml에서 tests=69, failures=0, errors=0, skipped=1(xfail). 합성 가중치 50→60 조작 시 결과 비율 변경도 브라우저에서 확인.",55)

# 10 — paced future reports rather than a premature final presentation.
s=page(10,"09  정리와 다음 단계","이번 단계는 ‘자료 확보 + 구현 기반’까지","남은 3~4회의 발표에서는 자료 연결, 사례 적용, 비교·검증을 차례로 보여줄 계획이다.")
summary=[("확보","도로망·경계·교통량·속도",TEAL,MINT),("구현","분석 엔진·가중치·웹 화면",BLUE,PALE),("다음 과제","인증·자료 변환·실제 적용",AMBER,CREAM)]
for i,(t,d,c,fill) in enumerate(summary):
    x=.71+i*4.97;box(s,x,2.30,4.65,1.03,fill,radius=True)
    text(s,x+.22,2.47,4.20,.33,t,15,c,True);text(s,x+.22,2.89,4.20,.32,d,17,INK,True)
roadmap=[("4차","자료 연결","토큰 발급·인구 파일 확보\n공식 연평균 자료 변환"),("5차","사례 적용","경로·동기간 시설정보 보강\n마포구 후보 분석 연결"),("6차","비교·검증","모델·가중치 비교\n조건 변화와 순위 안정성"),("7차 / 선택","최종 통합","전체 실행 시연\n보고서·한계·후속 과제")]
for i,(num,t,desc) in enumerate(roadmap):
    x=.72+i*3.72
    circle(s,x+.07,3.95,.19,BLUE if i<3 else TEAL)
    if i<3:line(s,x+.3,4.045,x+3.65,4.045,LINE,2)
    text(s,x+.03,4.44,3.42,.29,num,14,BLUE if i<3 else TEAL,True)
    text(s,x+.03,4.96,3.43,.45,t,25,INK,True)
    text(s,x+.03,5.68,3.43,.86,desc,16,MUTED)
box(s,.72,6.98,14.56,.73,NAVY,radius=True)
text(s,.95,7.19,14.05,.35,"다음 목표: 확보한 자료를 의미에 맞게 연결하고, 실제 입력으로 작동하는 사례를 제시한다.",20,WHITE,True,PP_ALIGN.CENTER)
source(s,"발표 횟수·승인 일정에 따라 마지막 두 단계를 통합 가능 | 실제 배송 GPS·독립 평가 자료의 미확보 상태는 계속 명시")
notes(s,"10. 다음 발표로 이어갈 계획", "이번 단계의 성과는 네 종류의 공식 파일을 확보하고, 이 자료를 연결할 입력 조건을 확인했으며, 분석 기능과 사용자 화면을 구현했다는 점입니다. 다음 발표에서는 토큰과 인구 파일을 확보하고 공식 연평균 자료를 처리하는 변환 기능을 연결하겠습니다. 그다음에는 경로와 기간이 맞는 시설정보를 보강해 마포구 사례를 적용하고, 이후 모델과 가중치에 따라 후보 순위가 얼마나 달라지고 유지되는지 검토하겠습니다. 실제 배송 GPS가 없으면 가상 순회 시나리오라는 점을 명시하겠습니다. 발표가 세 번 남으면 비교·검증과 최종 시연을 묶고, 네 번 남으면 나누어 진행할 수 있습니다. 이상으로 3차 진행 보고를 마치겠습니다.", "미래 계획이며 수행 결과가 아님. 인구 CSV 토큰 발급 및 데이터 제공 범위에 따라 일정 조정. 일별 원자료 없이 연평균 값을 복제해 일 단위 Bootstrap을 수행하지 않음. 순위 변화만으로 모델 성능 향상을 주장하지 않음.",55)

assert len(prs.slides)==10
path=OUT/"Transit_OOH_Capstone_3rd.pptx";prs.save(path)
(OUT/"slide_content.json").write_text(json.dumps(slides,ensure_ascii=False,indent=2),encoding="utf-8")

# Speaker notes: accessible, separate from the fixed 10-slide deck.
doc=Document();section=doc.sections[0]
section.top_margin=section.bottom_margin=DInches(.70);section.left_margin=section.right_margin=DInches(.78)
for style_name in ["Normal","Title","Heading 1","Heading 2"]:
    st=doc.styles[style_name];st.font.name=FONT;st._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),FONT)
doc.styles["Normal"].font.size=DPt(10.5)
doc.styles["Normal"].paragraph_format.space_after=DPt(7)
doc.styles["Normal"].paragraph_format.line_spacing=1.18
doc.styles["Title"].font.color.rgb=DColor.from_string(NAVY)
doc.styles["Heading 1"].font.color.rgb=DColor.from_string(BLUE)
doc.add_heading("Transit OOH\n3차 발표 대본 및 참고 메모",0)
doc.add_paragraph("정보통신공학과 · 캡스톤디자인2 | 10장 · 약 9분 | 자료 확인 기준 2026.10.11")
doc.add_paragraph("이번 발표의 중심: 데이터를 어떻게 확보했고, 어디까지 확인했으며, 시스템을 어떻게 구현했는가. 실제 자료와 합성 데모, 완료와 계획을 구분해 설명한다.")
for i,info in enumerate(slides):
    doc.add_heading(info["title"],1)
    doc.add_paragraph(f"권장 발표시간: 약 {info['seconds']}초")
    doc.add_paragraph(info["script"])
    p=doc.add_paragraph("발표자 참고: "+info["evidence"])
    for r in p.runs:r.font.size=DPt(9);r.font.color.rgb=DColor.from_string(MUTED)

doc.add_page_break();doc.add_heading("질문이 나왔을 때",1)
qas=[
 ("실제 데이터를 확보했나요?","네. View-T의 도로망, 경계, 추정교통량, 평균속도 원본을 받았습니다. 다만 인구 CSV와 실제 배송 GPS는 아직 없고, 대시보드 시연은 합성 입력입니다."),
 ("왜 마포구인가요?","전국 분석에 앞서 한 자치구 안에서 자료의 연결과 시스템 작동을 확인하기 위한 예시 지역입니다. 마포구가 최적 지역이라고 검증한 것은 아닙니다."),
 ("408은 차량 수인가요?","아닙니다. 공급자가 마포구로 분류한 주요도로 구간 수입니다. 교통량·속도·도로망에서 이 도로 번호 408개가 일치했습니다."),
 ("같은 16인데, 지도 경계와 앱 후보가 같은 건가요?","아닙니다. 확보한 경계는 마포구 실제 16개 동입니다. 앱 시연의 16개 후보는 합성 거점이며 그 동들을 분석한 결과가 아닙니다."),
 ("토큰만 받으면 분석이 끝나나요?","파일 확보 뒤에도 공간 단위, 기간, 집계 의미를 확인하고 변환 기능을 연결해야 합니다. 토큰 발급과 분석 완성은 별개의 단계입니다."),
 ("정차시간을 교통량·속도로 계산할 수 있나요?","일반 교통의 평균속도로 개별 배송차량의 정차시간을 직접 측정할 수 없습니다. 해당 지표는 GPS가 있을 때 계산하고, 없을 때는 제외하거나 가정으로 표시합니다."),
 ("50:50의 근거는 무엇인가요?","두 정보를 같은 비중으로 비교하기 위한 기본값입니다. 기존 60:40도 초기 가정이며, 사용 목적과 데이터에 따라 바꾸고 결과의 변화를 확인합니다."),
 ("정보통신공학적으로 구현한 것은 무엇인가요?","시간·좌표 데이터의 오류 검사, 정차·방문 정보 추출, 공간 번호 기반 자료 연결, 점수 계산, 사용자 조작에 따른 결과 재계산, 웹 GIS와 파일 입출력입니다."),
 ("현재 시험 결과는 무엇을 증명하나요?","68개 자동검사 통과는 정의한 입력·계산·화면 동작이 작동하는지 확인한 결과입니다. 실제 광고 효과나 최적 입지를 증명하는 수치는 아닙니다."),
 ("연평균 자료로 선정 안정성을 시험할 수 있나요?","가중치를 바꿨을 때의 순위 변화는 검토할 수 있습니다. 날짜를 다시 뽑는 방식의 안정성 분석은 실제 일별 자료가 필요하며, 평균값을 여러 날짜에 복제해서 수행하지 않습니다."),
]
for q,a in qas:doc.add_heading(q,2);doc.add_paragraph(a)
doc.add_heading("주요 근거와 발표 시 주의할 구분",1)
for item in [
 "공식 제공처: 한국교통연구원 KTDB View-T — https://viewt.ktdb.go.kr/cong/map/page.do",
 "수집 상태: docs/VIEWT_COLLECTION_STATUS.md. 인구 CSV는 토큰 요구로 미확보, 공식 연평균 어댑터는 미완료.",
 "원본 파일: seoul_link_lev5.5_2024.zip / mapo_boundary_2023.zip / TrafficVolume(LINK).xlsx / AverageSpeed(LINK).xlsx.",
 "공간 그림: 공급자 마포구 할당 408개 도로 + 2023년 16개 동 경계. 경계 연도 동일성은 미확인. 지도는 참고 표시이며 동별 교통량 분석 결과가 아님.",
 "프로그램 근거: src/transit_ooh/scoring.py, viewt.py, app/streamlit_app.py, app/viewt_page.py. 실제 도로 경로 최적화 완료를 주장하지 않음.",
 "자동검사 근거: reports/runs/tests.xml, 2026-10-10 실행, 68 pass + 1 expected fail. 별도로 발표자료 제작 중 실제 브라우저의 가중치 변경 확인.",
 "인구 유입 화면에서 표출단위는 출발지역 분포. 시군구를 읍면동으로 바꾸는 것만으로 마포구 16개 도착 동별 유입이 산출되지는 않음.",
 "이 발표 자료는 연구 진행 보고용이다. 실제 관측 배송경로, 광고 노출 인원, 최적 가중치 또는 실측 성능 개선을 확보한 것으로 설명하지 않는다.",
]:doc.add_paragraph(item,style="Normal")
doc.save(OUT/"Transit_OOH_Capstone_3rd_Speaker_Notes.docx")
md=["# Transit OOH 3차 발표 대본", "", "10장 / 약 9분 / 자료 확인 기준 2026.10.11", ""]
for info in slides:md.extend([f"## {info['title']}","",info["script"],"",f"근거: {info['evidence']}",""])
(OUT/"발표대본.md").write_text("\n".join(md),encoding="utf-8")
print(f"Created {path.name}: {len(prs.slides)} slides, {sum(s['seconds'] for s in slides)} seconds of speaker notes.")

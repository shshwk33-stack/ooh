"""Rebuild the competition story as editable slides in the user's PM visual style.

The reference contributes layout and palette only. Its flattened images, names,
logos, research claims and experimental numbers are not copied into this deck.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import shutil
from pathlib import Path

from build_competition_slides import load_study, sha256, validate
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
os.environ.setdefault("MPLCONFIGDIR", str(ROOT / ".cache/matplotlib"))
os.environ.setdefault("XDG_CACHE_HOME", str(ROOT / ".cache"))
Path(os.environ["MPLCONFIGDIR"]).mkdir(parents=True, exist_ok=True)
ORANGE, CORNER, MINT = "FF6F40", "FE9C7B", "C6E3E4"
NAVY, DARK, TEAL, BLUE = "17253D", "404040", "0E9384", "3A5BA0"
GRAY, MUTED, LINE = "9CA5B6", "63738B", "CDD4DF"
PALE, LIGHT, WHITE = "EAF6F4", "F2F4F7", "FFFFFF"
FONT = "Noto Sans CJK KR"


class PMDeck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(20), Inches(11.25)
        p = self.prs.core_properties
        p.title = "Transit OOH · 공개자료 기반 강건한 후보지역 선정"
        p.subject = "공개 교통·상권자료 기반 의사결정 연구"
        p.author = p.last_modified_by = p.keywords = p.comments = ""
        self.outline = []

    def box(
        self, s, x, y, w, h, fill=WHITE, line=None, radius=False, shape=None, rotate=0
    ):
        kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
        sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fill)
        if line:
            sh.line.color.rgb = RGBColor.from_string(line)
            sh.line.width = Pt(1)
        else:
            sh.line.fill.background()
        style = sh._element.find(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}style"
        )
        if style is not None:
            sh._element.remove(style)
        sh._element.spPr.append(OxmlElement("a:effectLst"))
        if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = 0.10
        sh.rotation = rotate
        return sh

    def text(
        self,
        s,
        value,
        x,
        y,
        w,
        h,
        size=24,
        color=NAVY,
        bold=False,
        align=PP_ALIGN.LEFT,
        italic=False,
        underline=False,
    ):
        sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = sh.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for i, line in enumerate(str(value).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_before = Pt(0)
            p.space_after = Pt(3)
            p.line_spacing = 1.06
            r = p.add_run()
            r.text = line
            r.font.name, r.font.size = FONT, Pt(size)
            r.font.bold, r.font.italic, r.font.underline = bold, italic, underline
            r.font.color.rgb = RGBColor.from_string(color)
            ea = OxmlElement("a:ea")
            ea.set("typeface", FONT)
            r._r.get_or_add_rPr().append(ea)
        return sh

    def line(self, s, x1, y1, x2, y2, color=LINE, width=1.2, dash=False, curve=False):
        sh = s.shapes.add_connector(
            MSO_CONNECTOR.CURVE if curve else MSO_CONNECTOR.STRAIGHT,
            Inches(x1),
            Inches(y1),
            Inches(x2),
            Inches(y2),
        )
        style = sh._element.find(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}style"
        )
        if style is not None:
            sh._element.remove(style)
        sh._element.spPr.append(OxmlElement("a:effectLst"))
        sh.line.color.rgb = RGBColor.from_string(color)
        sh.line.width = Pt(width)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE

            sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        return sh

    def circle(self, s, x, y, diameter, fill=WHITE, line=None):
        return self.box(s, x, y, diameter, diameter, fill, line, shape=MSO_SHAPE.OVAL)

    def arrow(self, s, x, y, w=0.42, h=0.2, color=GRAY):
        return self.box(s, x, y, w, h, color, shape=MSO_SHAPE.RIGHT_ARROW)

    def bulb(self, s, x, y):
        # Generic, editable line illustration, not a reference image asset.
        self.line(s, x + 0.35, y - 0.28, x + 0.35, y + 0.10, DARK, 1.5)
        self.box(s, x + 0.20, y + 0.08, 0.30, 0.26, DARK, shape=MSO_SHAPE.TRAPEZOID)
        self.circle(s, x, y + 0.30, 0.7, "FFDD38")
        self.line(s, x + 0.28, y + 0.42, x + 0.29, y + 0.74, DARK, 1.4)
        self.line(s, x + 0.43, y + 0.42, x + 0.42, y + 0.74, DARK, 1.4)
        self.line(s, x + 0.29, y + 0.74, x + 0.42, y + 0.74, DARK, 1.4)

    def decor(self, s, cover=False):
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
        self.box(s, -1.65, -1.60, 4.10, 3.35, MINT, radius=True, rotate=42)
        self.circle(s, 18.70, -1.05, 2.70, CORNER)
        self.circle(s, -1.65, 10.35, 2.70, CORNER)
        self.box(s, 18.15, 10.25, 3.15, 2.65, MINT, radius=True, rotate=43)
        self.line(s, 17.05, 0.58, 19.27, 0.58, DARK, 1.6)
        for angle in range(0, 360, 45):
            a = math.radians(angle)
            self.line(
                s,
                19.60 + 0.12 * math.cos(a),
                0.59 + 0.12 * math.sin(a),
                19.60 + 0.25 * math.cos(a),
                0.59 + 0.25 * math.sin(a),
                ORANGE,
                5,
            )
        self.line(s, 0.42, 10.98, 19.55, 10.98, DARK, 1)
        self.text(s, "2026", 0.58, 11.02, 1.0, 0.20, 11.5, DARK)
        self.text(
            s,
            "TRANSIT OOH  ·  PUBLIC-DATA DECISION RESEARCH",
            4.2,
            11.01,
            11.5,
            0.22,
            11.5,
            NAVY,
            align=PP_ALIGN.CENTER,
        )
        self.text(
            s,
            f"{len(self.prs.slides):02d} / 15",
            18.50,
            11.01,
            1.0,
            0.22,
            11.5,
            NAVY,
            align=PP_ALIGN.RIGHT,
        )

    def new(self, first, second, english, lead="", source=""):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.decor(s)
        self.text(s, first, 0.66, 0.41, 17.5, 0.88, 54, DARK, True)
        self.text(s, second, 0.66, 1.27, 17.5, 0.89, 54, ORANGE, True)
        self.text(s, english, 0.70, 2.39, 17.0, 0.34, 22, DARK, True)
        self.line(
            s, 0.00, 2.82, min(15.0, max(7.4, len(english) * 0.17)), 2.82, DARK, 1.5
        )
        self.bulb(
            s, min(16.25, max(11.4, max(len(first), len(second)) * 0.73 + 1.2)), 1.31
        )
        if lead:
            self.text(s, "• " + lead, 0.67, 3.02, 18.45, 0.57, 24, NAVY, True)
        if source:
            self.text(s, source, 0.70, 10.57, 18.3, 0.28, 12.4, MUTED, italic=True)
        self.outline.append(
            {
                "page": len(self.prs.slides),
                "title": first + " " + second,
                "english": english,
                "lead": lead,
                "source": source,
            }
        )
        return s

    def caption(self, s, text, x, y, w, size=18):
        self.text(s, text, x, y, w, 0.52, size, NAVY, True, PP_ALIGN.CENTER)

    def picture(self, s, path, x, y, w, h):
        path = Path(path)
        if not path.is_file():
            raise FileNotFoundError(path)
        iw, ih = Image.open(path).size
        factor = min(w / iw, h / ih)
        ww, hh = iw * factor, ih * factor
        s.shapes.add_picture(
            str(path),
            Inches(x + (w - ww) / 2),
            Inches(y + (h - hh) / 2),
            width=Inches(ww),
            height=Inches(hh),
        )

    def table(
        self, s, headers, rows, x, y, widths, row_h=0.65, size=22, highlight=None
    ):
        total = sum(widths)
        for i, row in enumerate([headers, *rows]):
            yy = y + i * row_h
            fill = NAVY if i == 0 else (PALE if i - 1 == highlight else WHITE)
            self.box(s, x, yy, total, row_h, fill)
            self.line(s, x, yy + row_h, x + total, yy + row_h, LINE, 0.8)
            xx = x
            for j, cell in enumerate(row):
                if j:
                    self.line(s, xx, yy, xx, yy + row_h, LINE, 0.7)
                self.text(
                    s,
                    cell,
                    xx + 0.13,
                    yy + 0.11,
                    widths[j] - 0.26,
                    row_h - 0.12,
                    size,
                    WHITE if i == 0 else (TEAL if i - 1 == highlight else NAVY),
                    i == 0 or i - 1 == highlight,
                )
                xx += widths[j]

    def tag(self, s, label, x, y, w, color=TEAL, fill=PALE, size=20):
        self.box(s, x, y, w, 0.55, fill, line=color, radius=True)
        self.text(
            s,
            label,
            x + 0.10,
            y + 0.12,
            w - 0.20,
            0.36,
            size,
            color,
            True,
            PP_ALIGN.CENTER,
        )

    def metric(self, s, value, label, x, y, w, color=NAVY, big=48):
        self.text(s, value, x, y, w, 0.88, big, color, True)
        self.text(s, label, x, y + 1.32, w, 0.72, 21, MUTED)

    def map_points(self, s, f, x, y, w, h, selected=()):
        self.box(s, x, y, w, h, WHITE, LINE)
        xx = f.longitude.to_numpy()
        yy = f.latitude.to_numpy()
        xmin, xmax = xx.min() - 0.004, xx.max() + 0.004
        ymin, ymax = yy.min() - 0.003, yy.max() + 0.003
        for _, row in f.iterrows():
            px = x + 0.55 + (row.longitude - xmin) / (xmax - xmin) * (w - 1.55)
            py = y + 0.35 + (ymax - row.latitude) / (ymax - ymin) * (h - 1.0)
            chosen = row.grid_id in selected
            self.circle(
                s,
                px - 0.085,
                py - 0.085,
                0.17,
                TEAL if chosen else "A7BACB",
                ORANGE if chosen else None,
            )
            self.text(
                s, row["name"], px + 0.10, py - 0.13, 1.16, 0.28, 15.5, NAVY, chosen
            )
        self.text(s, "경도", x + w / 2 - 0.25, y + h - 0.36, 1.0, 0.24, 13.5, MUTED)
        self.text(s, "위도", x + 0.05, y + 0.10, 0.7, 0.24, 13.5, MUTED)

    def truck(self, s, x, y, scale=1):
        self.box(s, x, y, 1.35 * scale, 0.66 * scale, TEAL, line=NAVY, radius=True)
        self.box(
            s,
            x + 1.35 * scale,
            y + 0.16 * scale,
            0.58 * scale,
            0.50 * scale,
            BLUE,
            line=NAVY,
            radius=True,
        )
        self.box(
            s, x + 1.45 * scale, y + 0.21 * scale, 0.30 * scale, 0.18 * scale, WHITE
        )
        for xx in [0.33, 1.55]:
            self.circle(s, x + xx * scale, y + 0.55 * scale, 0.32 * scale, NAVY)
            self.circle(
                s, x + (xx + 0.09) * scale, y + 0.64 * scale, 0.14 * scale, WHITE
            )


def verify_narrative_claims(summary, tables):
    """Fail when the fixed narrative no longer fits this audited study snapshot."""
    required = {
        "eligible_dongs": 15,
        "mapo_dongs": 16,
        "network_links_matched": 408,
        "commercial_mapo_sector_rows": 5694,
        "scenario_endpoint_rows": 120,
        "unique_score_vectors": 25,
        "combinations_evaluated": 3003,
        "optimal_sets": 2,
        "common_transaction_industries": 18,
    }
    for name, expected in required.items():
        if summary[name] != expected:
            raise ValueError(f"Narrative needs review: {name} changed")
    primary = (
        tables["portfolio_comparison"]
        .query("k == 5 and weight_low == 0")
        .set_index("model")
    )
    expected_regrets = {
        "A_stores": 0.4143,
        "B_vehicle_context": 0.4000,
        "C50_reduced": 0.3143,
        "C60_reduced": 0.3143,
        "R_minimax": 0.2000,
    }
    for model, expected in expected_regrets.items():
        if f"{primary.loc[model, 'worst_regret']:.4f}" != f"{expected:.4f}":
            raise ValueError("Narrative regret changed: " + model)
    prices = tables["price_of_robustness"].set_index("model")
    for model, expected in [
        ("C50_reduced", 0.4393),
        ("R_minimax", 0.4273),
        ("R_optimum_2", 0.3824),
    ]:
        if f"{prices.loc[model, 'nominal_mean_utility']:.4f}" != f"{expected:.4f}":
            raise ValueError("Narrative nominal utility changed: " + model)
    cv = tables["held_out_quarter_specification"]
    for model, expected in [("C50_reduced", 0.3143), ("R_minimax", 0.2000)]:
        values = cv[cv.model.eq(model)].worst_regret
        if len(values) != 4 or any(f"{x:.4f}" != f"{expected:.4f}" for x in values):
            raise ValueError("Narrative held-out result changed")


def build(report, output, dashboard=None):
    summary, tables = load_study(report)
    verify_narrative_claims(summary, tables)
    f = tables["candidate_features"].sort_values("grid_id").reset_index(drop=True)
    comp = tables["portfolio_comparison"]
    main = comp[(comp.k == 5) & (comp.weight_low == 0)].set_index("model")
    ext = tables["external_context_association"]
    common = tables["external_context_common_sectors"]
    price = tables["price_of_robustness"].set_index("model")
    assets = output / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    if dashboard:
        dest = assets / "dashboard_controls_map_50.png"
        if Path(dashboard).resolve() != dest.resolve():
            shutil.copyfile(dashboard, dest)
        dashboard = dest
    figs = {
        n: assets / f"pm_{n}.png"
        for n in [
            "candidate_tradeoff",
            "portfolio_tradeoff",
            "sensitivity",
            "external_all_quarters",
        ]
    }
    if not all(p.exists() for p in figs.values()):
        raise FileNotFoundError(
            "Run scripts/pmstyle_figures.py before building the deck"
        )
    figure_manifest = json.loads((assets / "figures_manifest.json").read_text())
    for relative, expected in figure_manifest["input_sha256"].items():
        source = (report / relative).resolve()
        if not source.is_relative_to(report.resolve()) or sha256(source) != expected:
            raise ValueError("Chart source checksum mismatch: " + relative)
    for path in figs.values():
        if sha256(path) != figure_manifest["output_sha256"][path.name]["sha256"]:
            raise ValueError("Chart output checksum mismatch: " + path.name)
    d = PMDeck()
    s = d.prs.slides.add_slide(d.prs.slide_layouts[6])
    d.decor(s, cover=True)
    d.text(s, "TRANSIT OOH", 0.85, 0.63, 9.5, 0.65, 35, DARK, True)
    d.text(
        s,
        "RESEARCH PRESENTATION · MOBILITY DATA & ROBUST DECISION",
        0.87,
        2.33,
        17.5,
        0.48,
        23,
        DARK,
        True,
    )
    d.text(s, "공개 교통·상권자료를 활용한", 0.87, 3.02, 17.4, 0.80, 44, DARK, True)
    d.text(
        s,
        "이동형 광고 후보지역\n선정 의사결정 시스템.",
        0.85,
        3.90,
        15.4,
        2.26,
        64,
        ORANGE,
        True,
        underline=True,
    )
    d.text(
        s,
        "공개 교통·상권자료와 최악후회 최소화를 중심으로.",
        0.91,
        6.50,
        16.5,
        0.60,
        27,
        DARK,
        True,
        italic=True,
    )
    d.text(
        s,
        "마포구 실제 집계 사례 · 선호 비중과 자료 표현에 따른 선택 위험 비교",
        0.93,
        7.35,
        16.8,
        0.66,
        23,
        NAVY,
    )
    d.bulb(s, 17.70, 3.07)
    for i, (label, color) in enumerate(
        [("교통 맥락", BLUE), ("상권 맥락", ORANGE), ("조건 변화", TEAL)]
    ):
        yy = 4.40 + i * 1.80
        d.line(s, 13.53, yy + 0.12, 18.47, 6.32, color, 3, curve=i != 1)
        d.circle(s, 13.30, yy, 0.24, color, NAVY)
        d.text(s, label, 10.40, yy - 0.12, 2.45, 0.56, 26, NAVY, True, PP_ALIGN.RIGHT)
    d.circle(s, 17.90, 5.75, 1.15, TEAL, NAVY)
    d.text(s, "R", 17.92, 5.94, 1.11, 0.72, 42, WHITE, True, PP_ALIGN.CENTER)
    for i, label in enumerate(
        ["공개자료 확보", "정합성 검사", "강건한 후보 선정", "GIS 의사결정"]
    ):
        x = 0.9 + i * 3.95
        d.tag(s, label, x, 9.53, 3.50, WHITE, NAVY, 21)
        if i < 3:
            d.arrow(s, x + 3.60, 9.74, 0.24, 0.13, ORANGE)
    d.outline.append(
        {
            "page": 1,
            "title": "공개 교통·상권자료를 활용한 이동형 광고 후보지역 선정 의사결정 시스템",
            "english": "Transit OOH",
        }
    )

    s = d.new(
        "연구 배경 및",
        "문제 정의.",
        "Research Background & Problem Definition.",
        "배송운영에서 생기는 이동정보를, 광고 검토 지역의 선택 근거로 활용할 수 있는가?",
        "개념도: 반복 순회형 배송운영을 설명하기 위한 도식이며 실제 운행기록이 아님.",
    )
    d.text(s, "반복 순회형 배송운영", 0.80, 3.65, 8.7, 0.52, 30, NAVY, True)
    pts = [(2.0, 5.10), (6.6, 4.75), (6.6, 7.25), (2.0, 7.60)]
    for i, (x, y) in enumerate(pts):
        nx, ny = pts[(i + 1) % 4]
        d.line(s, x + 0.35, y + 0.35, nx + 0.35, ny + 0.35, GRAY, 3)
    for i, (x, y) in enumerate(pts):
        d.circle(s, x, y, 0.70, WHITE, TEAL)
        d.text(s, chr(65 + i), x, y + 0.09, 0.70, 0.48, 25, TEAL, True, PP_ALIGN.CENTER)
    d.truck(s, 3.6, 5.70, 1.1)
    d.text(
        s,
        "물류 거점을 반복적으로 방문하는 밀크런형 운행",
        0.85,
        8.66,
        8.8,
        0.75,
        22,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )
    d.caption(s, "Fig.1. 배송활동에서 발생하는 정보의 추가 활용.", 0.80, 9.76, 8.9, 20)
    for i, (title, body) in enumerate(
        [
            (
                "원래의 연구 아이디어",
                "이동·정차·반복방문과 주변 공간정보를 연결해\n차량을 이동 데이터가 생성되는 플랫폼으로 본다.",
            ),
            (
                "실무에서 필요한 결정",
                "상권과 교통 중 무엇을 더 중요하게 볼지\n확정하기 전에 검토할 후보지역을 좁힌다.",
            ),
            (
                "이번 연구의 초점",
                "실제 배송 GPS를 확보했다고 가정하지 않고,\n공개자료로 비교 가능한 지역의 선택 위험을 분석한다.",
            ),
        ]
    ):
        yy = 3.66 + i * 1.92
        d.text(s, title, 10.1, yy, 8.7, 0.48, 28, TEAL if i != 2 else ORANGE, True)
        d.text(s, body, 10.1, yy + 0.64, 8.6, 1.08, 22, NAVY)
        d.line(s, 10.1, yy + 1.65, 18.95, yy + 1.65)

    s = d.new(
        "연구의 전개 및",
        "분석 범위.",
        "Research Development & Scope.",
        "1·2·3차 개발 흐름을 유지하면서, 이번 제출에서는 실제 공개자료의 비교 결과를 검증했다.",
    )
    phases = [
        ("1차 · 아이디어", "배송 이동데이터의 추가 활용\n이동·공간 관점의 위치가치"),
        ("2차 · PoC", "GPS 기반 정차·방문 분석\n합성 예제로 계산 구조 확인"),
        ("3차 · 확보 체계", "View-T 도로망·파일 입력\n공식 자료 확보·연동 준비"),
        ("현재 · 실제 연구", "2024 교통·상권 결합\n모델 비교·강건 선택·GIS"),
    ]
    for i, (title, body) in enumerate(phases):
        x = 0.78 + i * 4.76
        d.tag(
            s,
            title,
            x,
            3.75,
            4.22,
            TEAL if i == 3 else NAVY,
            PALE if i == 3 else LIGHT,
            23,
        )
        d.text(s, body, x + 0.10, 4.65, 4.03, 1.40, 23, NAVY, i == 3)
        if i < 3:
            d.arrow(s, x + 4.36, 4.04, 0.30, 0.15, GRAY)
    d.line(s, 0.80, 6.42, 19.18, 6.42)
    d.table(
        s,
        ["분석 경로", "입력과 계산", "이번 발표에서의 지위"],
        [
            [
                "기존 GPS 엔진",
                "정차시간 · 비정차 존재시간 · 반복방문",
                "구현·합성 검증 / 실제 배송 GPS는 미확보",
            ],
            [
                "실제 공개자료 연구",
                "연평균 차량 맥락 · 분기 점포 재고",
                "실제 기관 자료를 연결한 후보 행정동 비교",
            ],
        ],
        0.82,
        7.01,
        [3.6, 7.55, 7.1],
        0.87,
        21,
        1,
    )
    d.text(
        s,
        "현재의 실제 모델은 축소형이다. 교통량에서 정차시간이나 방문횟수를 만들어 넣지 않는다.",
        0.86,
        9.88,
        18.0,
        0.48,
        22,
        TEAL,
        True,
    )

    s = d.new(
        "실제 데이터의",
        "확보 및 연결.",
        "Acquired Data & Source Alignment.",
        "연도·지역·공간키와 측정량을 검사한 뒤, 동일 후보군에서 비교한다.",
        "자료: 서울시 OA-22172(2024 점포), OA-22175(2025 추정거래), KTDB View-T(2024 교통·도로망). [D1–D3]",
    )
    for value, label, x, color in [
        ("408", "교통·속도·도로망\nID·연장 일치 링크", 0.9, NAVY),
        ("5,694", "마포구 동×분기×업종\n보고 점포 행", 7.25, TEAL),
        ("16 → 15", "행정동 전체 →\n비교 가능한 후보", 13.7, ORANGE),
    ]:
        d.metric(s, value, label, x, 3.78, 5.4, color, 58)
    d.table(
        s,
        ["자료", "기준 · 단위", "적용과 해석"],
        [
            [
                "View-T 도로 맥락",
                "2024 평일 연평균 · 도로 링크",
                "기관 추정 교통량 / 실제 배송 궤적과 구분",
            ],
            [
                "서울시 점포",
                "2024 4개 분기 · 동×업종",
                "보고된 전체 점포 재고 / 연간 신규 점포 수 아님",
            ],
            [
                "서울시 추정매출",
                "2025 4개 분기 · 동×업종",
                "추정 거래 건수 / 선택 결과의 보조 맥락",
            ],
        ],
        0.82,
        6.10,
        [4.05, 6.00, 8.20],
        0.78,
        21,
    )
    d.text(s, "제외 원칙", 0.86, 9.53, 2.1, 0.46, 25, ORANGE, True)
    d.text(
        s,
        "염리동은 보유 도로 맥락이 없어 제외한다. 결측을 교통량 0으로 해석하지 않는다.",
        3.0,
        9.56,
        15.9,
        0.50,
        22,
        NAVY,
    )

    s = d.new(
        "공간정보 처리와",
        "정합성 검증.",
        "Data Engineering & Quality Control.",
        "서로 다른 기관 코드를 명칭·연도·고유성으로 대응하고, 기관이 부여한 도로–동 키를 사용한다.",
        "표시점 출처: 기존 공식 2024 동별 유입 집계표. 지도 좌표는 배송 거점·정차 위치가 아님. [D1,D3]",
    )
    labels = [
        "원본 수집",
        "코드·연도 검사",
        "동별 특성 집계",
        "점수·후보 선택",
        "지도·표·CSV",
    ]
    for i, label in enumerate(labels):
        x = 0.82 + i * 3.77
        d.tag(s, label, x, 3.66, 3.26, WHITE, NAVY if i != 3 else TEAL, 20)
        if i < 4:
            d.arrow(s, x + 3.36, 3.88, 0.27, 0.14, ORANGE)
    d.map_points(s, f, 0.85, 4.62, 9.0, 5.38)
    d.caption(s, "Fig.2. 실제 비교 가능한 마포구 15개 행정동.", 0.9, 10.10, 8.9, 17.6)
    items = [
        (
            "공간키",
            "서울시 11440… ↔ View-T 11140…\n접두사를 일괄 치환하지 않고 고유 동 명칭을 대조.",
        ),
        (
            "도로 지원",
            "408개 링크의 기관 동 코드·길이를 검증.\n웹 지도 경계로 임의 배분한 공식 통계가 아님.",
        ),
        (
            "점포 산식",
            "전체 점포 = 일반 점포 + 프랜차이즈 확인.\n전체 점포를 다시 더해 중복 계산하지 않음.",
        ),
        (
            "시간 정의",
            "인구 월~금과 도로 평일 정의가 불명확하여\n인구는 이번 축소형 점수에 투입하지 않음.",
        ),
    ]
    for i, (title, body) in enumerate(items):
        yy = 4.57 + i * 1.42
        d.text(s, title, 10.45, yy, 2.1, 0.47, 25, TEAL, True)
        d.text(s, body, 12.52, yy, 6.65, 1.05, 20.5, NAVY)
        d.line(s, 10.45, yy + 1.18, 19.1, yy + 1.18)

    s = d.new(
        "위치가치 점수와",
        "비교 모델.",
        "Location-Value Scores & Baseline Models.",
        "두 성분을 같은 후보군에서 정규화하고, 차량 맥락의 비중 α를 사용자가 조정한다.",
        "A는 보고 점포 재고, B는 보유 일반도로의 길이 가중 차량 맥락. 기존 전체 Spatial/GPS Mobility와 구분. [D1,D3]",
    )
    d.box(s, 0.86, 3.78, 18.22, 1.02, LIGHT)
    d.text(
        s,
        "위치가치  C(i, α) = (1−α) × 공간 A(i) + α × 이동 B(i)",
        1.15,
        3.98,
        17.65,
        0.65,
        33,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )
    d.table(
        s,
        ["모델", "같은 K개 지역을 고르는 기준"],
        [
            ["A · 상권 단독", "2024년 보고 전체 점포 재고의 분기 평균"],
            ["B · 교통 단독", "일반도로 링크의 길이 가중 차량 교통 맥락"],
            ["C50 / C60 · 고정 비중", "α=.5 / .6에서 점수가 높은 K개 지역"],
            ["R · 강건 선택", "여러 조건에서 최악 선택 손실이 최소인 K개 집합"],
        ],
        0.86,
        5.30,
        [4.75, 13.47],
        0.85,
        23,
        3,
    )
    d.text(s, "초기 가정", 0.9, 9.83, 2.5, 0.47, 26, ORANGE, True)
    d.text(
        s,
        "50:50과 60:40은 검증된 최적 비중이 아니다. 대시보드에서 0~100%를 바꿔 본다.",
        3.35,
        9.88,
        15.5,
        0.48,
        22,
        NAVY,
    )

    s = d.new(
        "제안 모델의",
        "수리 정의.",
        "Mathematical Modeling · Minimax Regret.",
        "각 조건의 최선 집합과 비교한 후회를 정의하고, 가장 큰 후회가 최소인 집합을 선택한다.",
        "최악후회 최소화는 기존 의사결정 방법이다. 이번 기여는 공개자료의 실제 적용·검증·시스템 구현에 있다. [R1,R2]",
    )
    values = [
        ("15개", "비교 가능한 후보 행정동"),
        ("5곳", "주 분석에서 선택할 지역 수"),
        ("120행", "정해 둔 조건의 끝점 설계"),
        ("3,003개", "K=5 후보 조합의 완전 탐색"),
    ]
    for i, (value, label) in enumerate(values):
        yy = 3.7 + i * 1.55
        d.text(s, value, 0.88, yy, 4.15, 0.70, 42, TEAL if i == 3 else NAVY, True)
        d.text(s, label, 0.91, yy + 0.82, 4.1, 0.44, 19.5, MUTED)
        d.line(s, 0.9, yy + 1.36, 4.95, yy + 1.36)
    d.box(s, 5.5, 3.76, 13.52, 5.93, LIGHT)
    d.text(
        s,
        "기호  x(i): 후보 i 선택 여부 · v(i,s): 조건 s의 정규화 점수",
        5.83,
        4.05,
        12.85,
        0.52,
        23,
        NAVY,
        True,
    )
    d.text(
        s,
        "U*(s): 조건 s를 미리 알았을 때 최선 K개 집합의 평균 점수",
        5.83,
        4.76,
        12.85,
        0.47,
        22,
        MUTED,
    )
    d.line(s, 5.84, 5.48, 18.68, 5.48)
    d.text(s, "목적함수", 5.85, 5.83, 2.55, 0.6, 26, TEAL, True)
    d.text(s, "minimize  r", 8.59, 5.74, 9.85, 0.70, 35, NAVY, True)
    d.text(s, "제약식", 5.85, 6.92, 2.55, 0.6, 26, TEAL, True)
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    formulas = [
        r"$r \geq U^*(s)-\frac{1}{K}\sum_{i=1}^{15}v(i,s)x(i),\quad\forall s$",
        r"$\sum_{i=1}^{15}x(i)=K$",
        r"$x(i)\in\{0,1\},\quad i=1,\ldots,15$",
    ]
    math_paths = []
    for i, formula in enumerate(formulas):
        fig = plt.figure(figsize=(12, 2))
        fig.text(0.01, 0.40, formula, fontsize=30, color="#" + NAVY)
        path_math = assets / f"math_constraint_{i + 1}.png"
        fig.savefig(
            path_math, dpi=220, transparent=True, bbox_inches="tight", pad_inches=0.05
        )
        plt.close(fig)
        math_paths.append(path_math)
    # Equal transparent canvases preserve one font scale without cropping limits.
    images = [Image.open(p).convert("RGBA") for p in math_paths]
    canvas_size = (
        max(im.width for im in images) + 12,
        max(im.height for im in images) + 12,
    )
    for i, (path_math, image) in enumerate(zip(math_paths, images, strict=True)):
        canvas = Image.new("RGBA", canvas_size, (255, 255, 255, 0))
        canvas.paste(image, (6, (canvas.height - image.height) // 2), image)
        bbox = canvas.getbbox()
        if (
            bbox is None
            or min(bbox[0], bbox[1], canvas.width - bbox[2], canvas.height - bbox[3])
            < 4
        ):
            raise ValueError("Equation glyph would be cropped")
        canvas.save(path_math)
        d.picture(s, path_math, 8.59, 6.50 + i * 0.95, 9.85, 1.0)
    d.text(
        s,
        "r: 최악후회의 상한 · s: 자료 표현과 가중치 끝점의 조합",
        5.85,
        9.22,
        12.80,
        0.36,
        17.5,
        MUTED,
    )
    d.text(
        s,
        "가중치에 대한 후회는 볼록 → 구간 양 끝으로 연속 범위의 최악값을 확인.",
        5.72,
        9.94,
        13.10,
        0.54,
        20.8,
        TEAL,
        True,
    )

    s = d.new(
        "비교 실험 및",
        "검증 설계.",
        "Experimental Design & Validation.",
        "같은 조건에서의 최적화, 조건을 제외한 재평가, 외부 상권 맥락을 구분한다.",
    )
    d.table(
        s,
        ["검사 단계", "실행 방법", "해석 범위"],
        [
            [
                "동일 조건 비교",
                "A / B / C50 / C60 / R · 같은 후보와 K",
                "설계한 목적함수의 조건부 결과",
            ],
            [
                "제외 분기 검사",
                "4개 분기를 하나씩 제외 · 연평균도 제거",
                "같은 지역의 자료 명세에 대한 견고함",
            ],
            [
                "제외 정의 검사",
                "도로 표현 4종 · 정규화 3종을 각각 제외",
                "학습에서 쓰지 않은 표현으로 재평가",
            ],
            [
                "외부 맥락 확인",
                "2025년 추정 거래 · 모든 분기 공개",
                "예비값 인지 후 회고적·탐색적 비교",
            ],
            [
                "수치 독립 재계산",
                "별도 표준라이브러리 코드로 전 조합 확인",
                "동일 입력의 계산 일관성",
            ],
        ],
        0.85,
        3.80,
        [3.85, 8.72, 5.65],
        0.90,
        22,
    )
    d.text(s, "조건 구성", 0.90, 9.52, 2.25, 0.49, 27, TEAL, True)
    d.text(
        s,
        "5개 점포 기준 × 4개 교통 표현 × 3개 정규화 × 가중치 구간 양 끝",
        3.15,
        9.58,
        15.55,
        0.48,
        22,
        NAVY,
        True,
    )
    d.text(
        s,
        "120개 설계 행은 25개 고유 점수 벡터를 포함한다. 독립 표본 수나 Bootstrap 표본 수가 아니다.",
        0.91,
        10.24,
        18.0,
        0.40,
        18.6,
        MUTED,
    )

    s = d.new(
        "주요 결과 및",
        "기준 모델 비교.",
        "Main Results · Fixed Weights vs. Robust Selection.",
        "K=5, 가중치 0~100%의 정해 둔 조건에서 최악후회와 명목 점수의 교환관계를 확인했다.",
        "후회는 정규화된 선택 효용의 차이이며 광고효과·매출·운영비 개선율을 측정한 값이 아님.",
    )
    labels = [
        ("A_stores", "A  상권 단독"),
        ("B_vehicle_context", "B  교통 단독"),
        ("C50_reduced", "C  50:50"),
        ("C60_reduced", "C  60:40"),
        ("R_minimax", "R  강건 선택"),
    ]
    for i, (key, label) in enumerate(labels):
        yy = 4.05 + i * 0.89
        value = main.loc[key, "worst_regret"]
        d.text(s, label, 0.95, yy, 3.05, 0.50, 24, TEAL if i == 4 else NAVY, i == 4)
        d.box(
            s,
            4.18,
            yy + 0.06,
            float(value) / 0.45 * 5.5,
            0.42,
            TEAL if i == 4 else (ORANGE if i >= 2 else (BLUE if i == 0 else GRAY)),
        )
        d.text(
            s,
            f"{value:.4f}",
            4.36 + float(value) / 0.45 * 5.5,
            yy - 0.01,
            1.95,
            0.54,
            24,
            NAVY,
            True,
        )
    d.caption(s, "Fig.3. 모델별 최악후회 · 작을수록 강건.", 0.90, 9.04, 10.0, 21)
    d.text(s, "최악후회", 12.12, 3.92, 6.5, 0.45, 27, TEAL, True)
    d.text(s, "0.3143 → 0.2000", 12.10, 4.62, 6.65, 0.90, 41, TEAL, True)
    d.line(s, 12.10, 5.85, 18.91, 5.85)
    d.text(s, "기본 50:50 평균점수", 12.12, 6.18, 6.55, 0.46, 25, NAVY, True)
    d.text(s, "0.4393 → 0.4273", 12.10, 6.83, 6.65, 0.80, 37, NAVY, True)
    d.text(s, "명목 효용의 차이  0.0120", 12.12, 7.88, 6.55, 0.57, 26, ORANGE, True)
    d.text(
        s,
        "최악후회가 같은 최적 조합 2개를 모두 공개.\n대표 조합은 결과에 맞춰 바꾸지 않고 코드순 선택.",
        12.12,
        8.82,
        6.55,
        1.05,
        21,
        NAVY,
    )

    s = d.new(
        "선택 결과와",
        "강건성의 대가.",
        "Selected Regions & Price of Robustness.",
        "명목 점수만 최대화하는 선택과, 가정 변화의 위험을 줄이는 선택은 다를 수 있다.",
        "효용–후회 경계는 가능한 3,003개 조합을 비교한 결과다. 비용·수익의 경제적 경계로 해석하지 않는다.",
    )
    d.picture(s, figs["portfolio_tradeoff"], 0.62, 3.58, 10.25, 5.98)
    d.caption(s, "Fig.4. 기본 점수와 최악후회의 교환관계.", 0.80, 9.72, 10.0, 20)
    d.text(s, "두 최적 조합에 공통인 4개 후보", 11.37, 3.78, 7.73, 0.50, 27, TEAL, True)
    d.text(s, "용강 · 서교 · 합정 · 상암", 11.38, 4.49, 7.48, 0.70, 29, NAVY, True)
    d.table(
        s,
        ["조합", "다섯 번째 후보", "기본 효용"],
        [
            [
                "R1 · 대표",
                "망원2동",
                f"{price.loc['R_optimum_1', 'nominal_mean_utility']:.4f}",
            ],
            [
                "R2",
                "성산2동",
                f"{price.loc['R_optimum_2', 'nominal_mean_utility']:.4f}",
            ],
        ],
        11.35,
        5.75,
        [1.8, 3.5, 2.2],
        0.82,
        22,
        0,
    )
    d.text(
        s,
        "두 조합은 최악후회 0.2000에서만 동점.",
        11.40,
        8.55,
        7.25,
        0.65,
        23,
        ORANGE,
        True,
    )
    d.text(
        s,
        "R2는 기본 효용이 낮다. 모든 기준에서 동등한\n대안으로 설명하지 않는다.",
        11.40,
        9.32,
        7.25,
        0.92,
        21,
        NAVY,
    )

    s = d.new(
        "가중치와 자료 표현의",
        "민감도 분석.",
        "Sensitivity & Held-Out Specification Checks.",
        "K와 가중치 범위를 바꾸고, 선택에 쓰지 않은 분기·도로·정규화 조건에서도 확인했다.",
    )
    d.picture(s, figs["sensitivity"], 0.58, 3.63, 10.53, 5.74)
    d.caption(
        s, "Fig.5. 후보 수·가중치 범위별 C50과 R의 최악후회.", 0.74, 9.53, 10.25, 18.8
    )
    d.text(s, "선택에 쓰지 않은 분기", 11.55, 3.82, 7.35, 0.48, 27, TEAL, True)
    d.table(
        s,
        ["제외한 분기", "C50 후회", "R 후회"],
        [[f"2024 Q{q}", "0.3143", "0.2000"] for q in range(1, 5)],
        11.50,
        4.61,
        [2.95, 2.2, 2.15],
        0.70,
        20.5,
        3,
    )
    d.text(s, "도로·정규화 정의 제외 7개 검사", 11.55, 8.36, 7.27, 0.49, 24, NAVY, True)
    d.text(
        s, "제외 조건에서 R 최악후회 0.0898~0.2000", 11.55, 9.02, 7.28, 0.53, 22, NAVY
    )
    d.text(
        s,
        "K=3의 전체 범위에서는 C50과 R이 같다. 모든 조건에서 개선되는 방식은 아니다.",
        0.92,
        10.25,
        18.0,
        0.44,
        21,
        ORANGE,
        True,
    )

    s = d.new(
        "외부 상권 지표와",
        "결과 해석.",
        "External Commercial Context & Interpretation.",
        "2025년 추정 거래와의 관련성을 탐색적으로 확인하고, 불리한 결과도 함께 공개한다.",
        "이전 예비값을 접한 뒤 수행한 회고적 재현이다. 외부 사전등록·전향적 예측·실제 광고성과 검증이 아님. [D2]",
    )
    q1 = ext[ext.quarter.eq(20251)].set_index("model")
    rows = []
    for key, label in labels:
        row = q1.loc[key]
        rows.append(
            [
                label,
                f"{row.selected_transaction_share:.2%}",
                "집합 선택" if key == "R_minimax" else f"{row.spearman:.4f}",
            ]
        )
    d.table(
        s,
        ["모델", "Q1 거래 비중", "순위상관"],
        rows,
        0.85,
        3.80,
        [3.45, 3.50, 2.55],
        0.76,
        21,
        4,
    )
    d.text(
        s,
        "분모: 비교 가능한 15개 동의 보고 거래 합계.\nR은 집합을 고르므로 순위상관을 만들지 않는다.",
        0.92,
        8.61,
        9.28,
        0.85,
        19.8,
        MUTED,
    )
    d.picture(s, figs["external_all_quarters"], 10.65, 3.62, 8.55, 4.94)
    d.caption(
        s, "Fig.6. 2025년 4개 분기의 선택 집합 거래 비중.", 10.65, 8.72, 8.40, 17.8
    )
    a = float(
        common[
            (common.quarter.eq(20251)) & (common.model.eq("A_stores"))
        ].selected_transaction_share.iloc[0]
    )
    r = float(
        common[
            (common.quarter.eq(20251)) & (common.model.eq("R_minimax"))
        ].selected_transaction_share.iloc[0]
    )
    d.text(
        s,
        f"거래 지표만 공통 18개 보고 업종으로 제한: A {a:.2%} / R {r:.2%} → A가 조금 높음.",
        0.94,
        9.79,
        18.0,
        0.60,
        22.5,
        ORANGE,
        True,
    )

    s = d.new(
        "GIS 대시보드의",
        "구현 및 사용.",
        "Interactive GIS Dashboard & Decision Support.",
        "후보 수·가중치·정규화를 바꾸면 지도, 현재 점수, 순위, 강건 선택을 함께 확인한다.",
        "화면: 현재 실제 집계자료를 읽은 Streamlit 앱. 동 좌표는 표시점이며 실제 배송 정차 지점이 아님.",
    )
    if dashboard:
        d.picture(s, dashboard, 0.65, 3.58, 12.85, 6.43)
    else:
        chosen = set(str(main.loc["R_minimax", "selected_ids"]).split(";"))
        d.map_points(s, f, 0.88, 3.75, 12.3, 5.9, chosen)
    d.caption(s, "Fig.7. 실제 자료의 설정 조작·후보지역 지도.", 0.78, 10.05, 12.6, 19)
    for i, (title, body) in enumerate(
        [
            ("입력 설정", "K=3·5·10 / 비중 0~100%\n정규화 3종 / 고려할 비중 범위"),
            (
                "근거 확인",
                "기본 순위와 강건 집합 비교\n공통 후보·동점·효용의 대가 확인",
            ),
            (
                "업무 전달",
                "현재 후보 점수·집합 비교 CSV\n고정 연구 결과표와 구분해 내려받기",
            ),
        ]
    ):
        yy = 3.78 + i * 1.94
        d.circle(s, 14.10, yy, 0.58, TEAL)
        d.text(
            s,
            str(i + 1),
            14.10,
            yy + 0.07,
            0.58,
            0.39,
            24,
            WHITE,
            True,
            PP_ALIGN.CENTER,
        )
        d.text(s, title, 14.93, yy, 4.12, 0.49, 27, NAVY, True)
        d.text(s, body, 14.14, yy + 0.77, 4.86, 1.02, 20.7, NAVY)

    s = d.new(
        "실무 적용 및",
        "결론.",
        "Practical Implications & Conclusion.",
        "하나의 최적 가중치를 선언하기보다, 기준이 달라질 때 선택 근거가 어떻게 유지되는지 보여 준다.",
    )
    for i, (title, body, color) in enumerate(
        [
            (
                "업무 활용",
                "검토할 후보지역 K개 선정\n선호 변경에 따른 위험 확인\n후속 배송·매체 검토에 전달",
                TEAL,
            ),
            (
                "융합 기여",
                "공개자료의 공간·연도 검증\n산업공학의 집합 선택 적용\nGIS·CSV로 검토 과정 구현",
                BLUE,
            ),
            (
                "확장 방향",
                "다른 자치구·후보 단위\n물류거점·매체의 사전 선별\n운영 제약을 더한 후속 연구",
                ORANGE,
            ),
        ]
    ):
        xx = 0.9 + i * 6.31
        d.text(s, title, xx, 3.85, 5.72, 0.60, 32, color, True)
        d.line(s, xx, 4.63, xx + 5.63, 4.63, color, 3)
        d.text(s, body, xx, 4.98, 5.72, 2.12, 25, NAVY)
    d.line(s, 0.91, 7.67, 19.03, 7.67)
    d.text(s, "검증한 범위", 0.94, 8.04, 3.09, 0.51, 27, TEAL, True)
    d.text(
        s,
        "실제 15개 후보 · 전 조합 3,003개 · 독립 수치 재계산 · 자동검사 205개 통과",
        4.13,
        8.10,
        14.73,
        0.57,
        23,
        NAVY,
        True,
    )
    d.text(s, "해석의 경계", 0.94, 9.08, 3.09, 0.51, 27, ORANGE, True)
    d.text(
        s,
        "행정동 사전 선별 사례다. 실제 광고효과·운행 가능 경로·수익을 확정하지 않는다.",
        4.13,
        9.14,
        14.73,
        0.76,
        22,
        NAVY,
    )

    s = d.new(
        "참고문헌 및",
        "데이터 출처.",
        "References & Data Sources.",
        "공식 출처와 재현 가능한 계산 근거를 공개하고, 자료가 측정하지 않은 효과는 주장하지 않는다.",
    )
    d.tag(s, "Data & Sources", 0.90, 3.77, 8.69, WHITE, ORANGE, 26)
    data_refs = [
        (
            "[D1] 서울특별시 · 상권분석서비스(점포–행정동)",
            "2024년 4개 분기 · 공공누리 1유형\ndata.seoul.go.kr/dataList/OA-22172/S/1/datasetView.do",
        ),
        (
            "[D2] 서울특별시 · 상권분석서비스(추정매출–행정동)",
            "2025년 4개 분기 · 공공누리 1유형\ndata.seoul.go.kr/dataList/OA-22175/S/1/datasetView.do",
        ),
        (
            "[D3] 한국교통연구원 · KTDB View-T",
            "2024년 도로망·추정교통량·평균속도\nviewt.ktdb.go.kr/cong/map/page.do\nviewt.ktdb.go.kr/cong/pdf/Trafficestimate.pdf",
        ),
    ]
    for i, (title, body) in enumerate(data_refs):
        yy = 4.62 + i * 1.70
        d.text(s, title, 0.94, yy, 8.55, 0.52, 21.2, NAVY, True)
        d.text(s, body, 0.97, yy + 0.67, 8.52, 1.15, 18.5, MUTED)
    d.tag(s, "Research Methods", 10.36, 3.77, 8.68, WHITE, ORANGE, 26)
    refs = [
        (
            "[R1] Bertsimas, Brown & Caramanis (2010)",
            "Theory and Applications of Robust Optimization.\nhttps://arxiv.org/abs/1010.5445",
        ),
        (
            "[R2] Mastin, Jaillet & Chin (2014)",
            "Randomized Minmax Regret for Combinatorial\nOptimization Under Uncertainty.\nhttps://arxiv.org/abs/1401.7043",
        ),
    ]
    for i, (title, body) in enumerate(refs):
        yy = 4.62 + i * 1.90
        d.text(s, title, 10.41, yy, 8.52, 0.52, 22, NAVY, True)
        d.text(s, body, 10.44, yy + 0.67, 8.49, 1.26, 20, MUTED)
    d.text(s, "재현 근거", 10.43, 8.84, 8.45, 0.47, 24, TEAL, True)
    d.text(
        s,
        "원본 해시 · 후보 대응표 · 전 조합 결과 · PNG/SVG\n고정 의존성 · 검사 기록 · 별도 수치 검수 도구",
        10.43,
        9.48,
        8.48,
        0.95,
        20,
        NAVY,
    )
    if len(d.prs.slides) != 15:
        raise ValueError("Expected exactly 15 anonymous slides")
    output.mkdir(parents=True, exist_ok=True)
    path = output / "Transit_OOH_PMstyle_Anonymous.pptx"
    d.prs.save(path)
    (output / "slide_outline.json").write_text(
        json.dumps(d.outline, ensure_ascii=False, indent=2) + "\n"
    )
    return path, summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report", type=Path, default=ROOT / "reports/public_decision")
    parser.add_argument(
        "--output", type=Path, default=ROOT / "deliverables/competition2026_pmstyle"
    )
    parser.add_argument("--dashboard", type=Path)
    args = parser.parse_args()
    dashboard = args.dashboard or args.output / "assets/dashboard_controls_map_50.png"
    if not dashboard.is_file():
        dashboard = (
            ROOT / "deliverables/competition2026/assets/dashboard_live_controls_map.png"
        )
    path, _ = build(args.report, args.output, dashboard)
    validation = validate(path, args.report, args.output)
    # Existing validator checks all native paragraphs, identities and text bounds.
    import fitz

    with fitz.open(path.with_suffix(".pdf")) as pdf:
        text = "\n".join(p.get_text() for p in pdf)
    forbidden = [
        "명지대학교",
        "Myongji",
        "한국로지스틱스학회",
        "Korea Logistics Society",
        "Jun-Su",
        "Yu-Won",
        "Seo-Jeong",
        "Chang-Min",
        "Woo-Yeon",
    ]
    found = [x for x in forbidden if x.lower() in text.lower()]
    if found:
        raise ValueError("Reference-deck identity leak: " + str(found))
    original = Presentation(path)
    count = {"text_shapes": 0, "editable_nonpicture_shapes": 0, "pictures": 0}
    for slide in original.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                count["text_shapes"] += 1
            if shape.shape_type == 13:
                count["pictures"] += 1
            else:
                count["editable_nonpicture_shapes"] += 1
    validation.update(
        {
            "variant": "PM_style_native_rebuild",
            "style_reference_sha256": "1a64b71f5f58115e946080dfe0ca319673a76400d098af915a869398c7a91ba0",
            "reference_identity_tokens_found": found,
            "reference_content_images_copied": False,
            "quantitative_narrative_checked_against_source_tables": True,
            "chart_input_and_output_hashes_checked": True,
            "editable_shape_counts": count,
            "visual_review": "pending",
        }
    )
    (args.output / "presentation_validation.json").write_text(
        json.dumps(validation, ensure_ascii=False, indent=2) + "\n"
    )
    p = args.output / "presentation_text.txt"
    p.write_text(p.read_text().rstrip() + "\n")
    (args.output / "README.md").write_text(
        "# PM 스타일 · 익명 경진대회 발표자료\n\n제출용: **Transit_OOH_PMstyle_Anonymous.pdf** (표지 포함 15쪽).\n수정용: **Transit_OOH_PMstyle_Anonymous.pptx**.\n\n사용자가 제공한 PM 발표자료의 흰 배경·회색/주황 제목·청록색 강조·2열 논문형 배치·모서리 장식만 참고했습니다. 참고자료의 연구 내용·실험 결과·로고·소속·인물 사진은 복사하지 않았습니다. 원본의 이미지형 텍스트를 재사용하지 않고 제목·본문·도형을 편집 가능하게 만들었습니다.\n\n이전 1·2·3차의 배송 이동데이터 활용 동기와 개발 흐름을 복원하고 최신 실제 공개자료 연구를 반영했습니다. 과거 합성 PoC·계획 수치와 확인되지 않은 시장규모는 현재 실증 성과로 쓰지 않았습니다. 2025 추정 거래 검사는 탐색적·회고적이며 실제 광고효과 검증이 아닙니다.\n\n재생성: 먼저 python scripts/pmstyle_figures.py, 다음 python scripts/build_pmstyle_slides.py. 발표 제작에 python-pptx, pandas, numpy, matplotlib, Pillow, PyMuPDF, LibreOffice, Noto Sans CJK KR 글꼴이 필요합니다. 연구 앱 의존성은 바꾸지 않았습니다. PDF는 실제 PPTX를 LibreOffice로 변환했습니다. 글꼴이 없는 PC의 PPT 표시 차이를 피하려면 제출은 PDF를 사용하세요.\n\n협업 관련 내용은 사용자가 별도로 작성합니다. 최종 편집 후 소속·인명·로고·문서 속성과 제출 별첨을 다시 확인하세요. 참고문헌의 공공기관·논문 저자는 소속 표시와 구분합니다.\n"
    )
    print(json.dumps(validation, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

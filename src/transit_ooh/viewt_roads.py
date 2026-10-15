"""Official annual-average road files, kept separate from daily/GPS models."""

from __future__ import annotations

import hashlib
import json
import math
import re
import warnings
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ATTRIBUTES = [
    "link_id",
    "its_link_ids",
    "road_class",
    "road_name",
    "district",
    "length_km",
    "lanes",
]
METRICS = {
    "traffic": [
        "all_vehicles_per_day",
        "cars_per_day",
        "buses_per_day",
        "trucks_per_day",
    ],
    "speed": ["speed_kmh"],
}
TITLES = {"traffic": "추정교통량(도로구간)", "speed": "평균속도(도로구간)"}
COMMON_HEADERS = [
    "5.5LINKID",
    "ITSLINKID",
    "도로등급",
    "도로명",
    "권역",
    "연장(KM)",
    "차선수",
]


@dataclass
class RoadContext:
    table: pd.DataFrame
    metadata: dict


def _compact(value):
    return re.sub(r"\s+", "", str(value or ""))


def _identifier(value):
    if isinstance(value, bool) or value is None:
        raise ValueError("Missing or invalid road ID")
    if isinstance(value, str):
        result = value.strip()
        if not re.fullmatch(r"\d+", result):
            raise ValueError("Road IDs must be digit strings; no scientific notation")
        return result
    if not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError("Invalid numeric road ID")
    if value < 0 or value != int(value) or abs(value) >= 2**53:
        raise ValueError("Road IDs must be exact integers")
    return str(int(value))


def read_official_excel(source, kind):
    """Read the verified two-row provider layout, including wrong A1:A1 dimensions."""
    if kind not in TITLES:
        raise ValueError("Expected traffic or speed workbook")
    raw = source if isinstance(source, bytes) else Path(source).read_bytes()
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        try:
            book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        except (zipfile.BadZipFile, KeyError) as exc:
            raise ValueError("Invalid XLSX workbook") from exc
    try:
        if len(book.worksheets) != 1:
            raise ValueError("Expected one provider worksheet")
        sheet = book.worksheets[0]
        try:
            declared_dimension = sheet.calculate_dimension()
        except ValueError:
            declared_dimension = "unspecified"
        sheet.reset_dimensions()
        rows = []
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= 100_000:
                raise ValueError("Workbook exceeds supported 100000 rows")
            if any(v is not None for v in row):
                rows.append(tuple(row))
    finally:
        book.close()
    if not rows or _compact(rows[0][0]) != TITLES[kind]:
        raise ValueError(f"Expected official {TITLES[kind]} workbook")
    header_index = next(
        (
            i
            for i, row in enumerate(rows[:25])
            if _compact(row[0]).upper() == "5.5LINKID"
        ),
        None,
    )
    if header_index is None:
        raise ValueError("Road link header not found")
    preamble = [str(row[0]).strip() for row in rows[:header_index]]

    def field(label):
        found = [
            line.split(":", 1)[1].strip()
            for line in preamble
            if re.match(rf"^-\s*{label}\s*:", line)
        ]
        if len(found) != 1:
            raise ValueError(f"Missing or ambiguous metadata: {label}")
        return found[0]

    year = re.fullmatch(r"(\d{4})년", field("기준년도"))
    if not year:
        raise ValueError("Invalid reference year")
    day = field("기간")
    if day not in {"평일", "주말"}:
        raise ValueError("Unsupported day type")
    region = field("권역범위").split("/", 1)[0].strip()
    road_classes = sorted(v.strip() for v in field("도로등급").split(","))
    header = [_compact(v).upper() for v in rows[header_index]]
    expected_metrics = [
        f"{v}-{day}"
        for v in (["전체", "승용차", "버스", "트럭"] if kind == "traffic" else ["전체"])
    ]
    if header != COMMON_HEADERS + expected_metrics:
        raise ValueError(
            "Unsupported columns: require all classes for traffic and all vehicles for speed"
        )
    ncols = len(header)
    subheader = rows[header_index + 1] if header_index + 1 < len(rows) else ()
    if len(subheader) != ncols or any(_compact(v) != "연평균일" for v in subheader[7:]):
        raise ValueError(
            "Only annual-average daily values are supported; hourly data need a separate contract"
        )
    records = rows[header_index + 2 :]
    if not records or any(
        len(r) != ncols or any(v is None for v in r) for r in records
    ):
        raise ValueError("Empty table or missing cells / inconsistent column count")
    frame = pd.DataFrame(records, columns=ATTRIBUTES + METRICS[kind])
    frame["link_id"] = frame.link_id.map(_identifier)
    if frame.link_id.duplicated().any():
        raise ValueError("Duplicate road IDs")
    for name in ATTRIBUTES[1:5]:
        frame[name] = frame[name].astype(str).str.strip()
        if frame[name].eq("").any():
            raise ValueError(f"Missing road attribute: {name}")
    if not frame.district.eq(region).all():
        raise ValueError("District metadata and road rows disagree")
    for name in ["length_km", "lanes", *METRICS[kind]]:
        if frame[name].map(lambda v: isinstance(v, bool)).any():
            raise ValueError(f"Boolean numeric value: {name}")
        values = pd.to_numeric(
            frame[name].astype(str).str.replace(",", "", regex=False), errors="raise"
        )
        if not np.isfinite(values).all() or (values < 0).any():
            raise ValueError(f"Missing, negative or non-finite values: {name}")
        frame[name] = values
    if (
        frame.length_km.le(0).any()
        or frame.lanes.le(0).any()
        or frame.lanes.mod(1).ne(0).any()
    ):
        raise ValueError("Lengths and integral lane counts must be positive")
    if kind == "traffic":
        residual = frame.all_vehicles_per_day - frame[
            ["cars_per_day", "buses_per_day", "trucks_per_day"]
        ].sum(axis=1)
        if residual.abs().gt(2).any():
            raise ValueError("Vehicle class total differs beyond rounding tolerance")
    metadata = {
        "source": "KTDB View-T user-supplied provider-format download",
        "source_url": "https://viewt.ktdb.go.kr/cong/map/page.do",
        "kind": kind,
        "year": int(year[1]),
        "day_type": day,
        "district": region,
        "road_classes": road_classes,
        "network_level": "5.5",
        "statistic": "annual_average_day",
        "spatial_unit": "road_link",
        "unit": "vehicles/day" if kind == "traffic" else "km/h",
        "measurement_basis": "estimated" if kind == "traffic" else "provider_aggregate",
        "unit_basis": "official indicator definition; speed unit verified from provider screen",
        "raw_sha256": hashlib.sha256(raw).hexdigest(),
        "declared_dimensions": declared_dimension,
        "rows": len(frame),
        "preamble": preamble,
    }
    return RoadContext(frame.sort_values("link_id").reset_index(drop=True), metadata)


def combine_roads(traffic, speed):
    for key in [
        "year",
        "day_type",
        "district",
        "road_classes",
        "network_level",
        "statistic",
        "spatial_unit",
    ]:
        if traffic.metadata[key] != speed.metadata[key]:
            raise ValueError(f"Incompatible traffic/speed metadata: {key}")
    if set(traffic.table.link_id) != set(speed.table.link_id):
        raise ValueError("Traffic and speed road ID sets differ; no partial join")
    left = traffic.table.set_index("link_id")
    right = speed.table.set_index("link_id").reindex(left.index)
    for name in ATTRIBUTES[1:]:
        equal = (
            np.allclose(left[name], right[name], rtol=0, atol=1e-6)
            if name in {"length_km", "lanes"}
            else left[name].equals(right[name])
        )
        if not equal:
            raise ValueError(f"Road attributes disagree: {name}")
    result = left.join(right[["speed_kmh"]], validate="one_to_one").reset_index()
    result["year"] = traffic.metadata["year"]
    result["day_type"] = traffic.metadata["day_type"]
    result["statistic"] = "annual_average_day"
    result["traffic_basis"] = "estimated"
    result["speed_basis"] = "provider_aggregate"
    meta = {
        key: traffic.metadata[key]
        for key in [
            "year",
            "day_type",
            "district",
            "statistic",
            "network_level",
            "spatial_unit",
        ]
    }
    meta.update(
        {
            "schema": "viewt_official_road_context_v1",
            "rows": len(result),
            "sources": {"traffic": traffic.metadata, "speed": speed.metadata},
            "available_outputs": ["road_context_table", "descriptive_traffic_speed"],
            "unavailable_outputs": [
                "gps_stops",
                "repeat_visits",
                "population",
                "spatial_score",
                "integrated_score",
                "daily_bootstrap",
                "time_of_day_analysis",
            ],
            "limitations": [
                "Annual averages are not independent daily samples.",
                "All-vehicle speed is not delivery-truck speed or observed stop time.",
                "Link counts cannot be summed into district unique vehicles.",
                "Road context is not an OOH exposure count or a complete location-value model.",
            ],
        }
    )
    return RoadContext(result, meta)


def attach_network(context, archive, network_year):
    """Use native CRS and link keys; never attach a mismatched-year boundary."""
    import geopandas as gpd

    if network_year != context.metadata["year"]:
        raise ValueError("Network reference year must match traffic/speed")
    archive = Path(archive).resolve()
    with zipfile.ZipFile(archive) as z:
        shapes = [name for name in z.namelist() if name.lower().endswith(".shp")]
        if len(shapes) != 1:
            raise ValueError("Expected one road shapefile in ZIP")
        stem = shapes[0][:-4]
        if not all(stem + ext in z.namelist() for ext in [".shx", ".dbf", ".prj"]):
            raise ValueError("Road ZIP needs SHP, SHX, DBF and PRJ")
    network = gpd.read_file(f"zip://{archive}!{shapes[0]}", encoding="cp949")
    if network.crs is None or "k_link_id" not in network:
        raise ValueError("Network CRS and k_link_id are required")
    network["link_id"] = network.k_link_id.map(_identifier)
    if network.link_id.duplicated().any():
        raise ValueError("Duplicate network road IDs")
    selected = network[network.link_id.isin(context.table.link_id)].copy()
    if set(selected.link_id) != set(context.table.link_id):
        raise ValueError("Some traffic links have no network geometry")
    if (
        selected.geometry.isna().any()
        or selected.geometry.is_empty.any()
        or not selected.is_valid.all()
        or not selected.geom_type.isin(["LineString", "MultiLineString"]).all()
    ):
        raise ValueError("Invalid or unsupported road geometry")
    original_crs = selected.crs.to_wkt()
    mapped = (
        selected[["link_id", "geometry"]]
        .to_crs("EPSG:4326")
        .merge(context.table, on="link_id", validate="one_to_one")
    )
    collection = json.loads(mapped.to_json(drop_id=True))
    meta = {
        "network_year": network_year,
        "year_basis": "explicit input declaration; provider filename reviewed",
        "network_sha256": hashlib.sha256(archive.read_bytes()).hexdigest(),
        "original_crs_wkt": original_crs,
        "output_crs": "EPSG:4326",
        "matched_road_links": len(mapped),
        "boundary_overlay": "none",
        "semantics": "provider road links; not a tracked delivery route",
    }
    return collection, meta


def export_context(context, output, network=None, network_year=None):
    out = Path(output)
    # Resolve and validate before writing anything, including optional geometry.
    geometry = None
    metadata = dict(context.metadata)
    if network is not None:
        geometry, metadata["network"] = attach_network(context, network, network_year)
    out.mkdir(parents=True, exist_ok=True)
    csv = context.table.to_csv(index=False).encode("utf-8-sig")
    (out / "roads.csv").write_bytes(csv)
    metadata["table_sha256"] = hashlib.sha256(csv).hexdigest()
    metadata["geojson_sha256"] = None
    if geometry is not None:
        data = json.dumps(geometry, ensure_ascii=False, allow_nan=False).encode("utf-8")
        (out / "roads.geojson").write_bytes(data)
        metadata["geojson_sha256"] = hashlib.sha256(data).hexdigest()
    elif (out / "roads.geojson").exists():
        (out / "roads.geojson").unlink()
    (out / "metadata.json").write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2, allow_nan=False),
        encoding="utf-8",
    )
    return metadata


def load_context(directory):
    directory = Path(directory)
    metadata = json.loads((directory / "metadata.json").read_text())
    raw = (directory / "roads.csv").read_bytes()
    if (
        metadata.get("schema") != "viewt_official_road_context_v1"
        or hashlib.sha256(raw).hexdigest() != metadata["table_sha256"]
    ):
        raise ValueError("Road context schema or table checksum mismatch")
    table = pd.read_csv(BytesIO(raw), dtype={"link_id": str, "its_link_ids": str})
    if table.link_id.duplicated().any() or len(table) != metadata["rows"]:
        raise ValueError("Road context row counts or identifiers mismatch")
    geometry = None
    if metadata.get("geojson_sha256"):
        raw = (directory / "roads.geojson").read_bytes()
        if hashlib.sha256(raw).hexdigest() != metadata["geojson_sha256"]:
            raise ValueError("Road geometry checksum mismatch")
        geometry = json.loads(raw)
        if {f["properties"]["link_id"] for f in geometry["features"]} != set(
            table.link_id
        ):
            raise ValueError("Map and table road IDs differ")
    return RoadContext(table, metadata), geometry
